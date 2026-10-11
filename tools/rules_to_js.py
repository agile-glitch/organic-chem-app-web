"""Turn reaction_rules.xlsx + molecule_data.xlsx into data/reaction_rules.js for the Reactions page.

A web page cannot read an Excel file, so run this after editing either workbook:
    python tools/rules_to_js.py
"""
import json, os, re, sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
KEY = re.compile(r"[A-Z]{14}-[A-Z]{10}-[A-Z]")

def rows(path, sheet):
    ws = load_workbook(os.path.join(ROOT, path), read_only=True)[sheet]
    it = ws.iter_rows(values_only=True)
    head = [str(h) for h in next(it)]
    return [dict(zip(head, r)) for r in it if r and r[0]]

molecules = {}
for m in rows("molecule_data.xlsx", "Molecules"):
    molecules[m["InChIKey (ID)"]] = {"smiles": m["SMILES + stereo"], "name": m["Common name"] or m["IUPAC name"]}

def first_num(text):
    """the first number in a pKa cell such as '≈ 19–20 (α C–H)' or '≈ −6 (protonated C=O)'"""
    m = re.search(r"[-−]?\d+(?:\.\d+)?", str(text or "").replace("–", " "))
    return float(m.group(0).replace("−", "-")) if m else None

# the fragment library (sheet "Fragments"): the SMARTS the page uses to read the starting materials
fragments = {}
for fr in rows("molecule_data.xlsx", "Fragments"):
    if fr.get("SMARTS pattern"):
        fragments[str(fr["Fragment ID"]).strip()] = {"name": fr["Fragment name"], "smarts": fr["SMARTS pattern"],
                                                     "pKa": first_num(fr.get("pKa (its most acidic H)")),
                                                     "pKaH": first_num(fr.get("pKa of its conjugate acid (how basic)"))}
FID = re.compile(r"F\d{3}")
def frag_items(text):
    """'F018 aldehyde | F019 ketone + F017 1,2-diol' -> [["F018", "F019"], ["F017"]]: all items needed, any id within one"""
    return [ids for ids in (FID.findall(part) for part in str(text or "").split("+")) if ids]

def find(row, start):
    """the column whose header starts with `start` (headers carry notes in brackets)"""
    for k, v in row.items():
        if k.startswith(start):
            return v or ""
    return ""

TEMP_REQ = re.compile(r"^(>=|<=|≥|≤|>|<)?\s*(-?\d+(?:\.\d+)?)\s*(?:(?:-|–|to)\s*(-?\d+(?:\.\d+)?))?\s*°?\s*c$")
def requirement(text):
    """one required item: a named condition {"cond": "reflux"} or a temperature {"tmin", "tmax", "label"} (°C)"""
    s = text.strip().lower()
    m = TEMP_REQ.match(s)
    if not m:
        return {"cond": s}
    op, a, b = m.group(1), float(m.group(2)), m.group(3)
    fmt = lambda v: f"{v:g}".replace("-", "−")
    if b is not None:
        lo, hi = sorted((a, float(b)))
        return {"tmin": lo, "tmax": hi, "label": f"{fmt(lo)}–{fmt(hi)} °C"}
    if op in (">=", "≥", ">"):
        return {"tmin": a, "tmax": None, "label": f"≥ {fmt(a)} °C"}
    if op in ("<=", "≤", "<"):
        return {"tmin": None, "tmax": a, "label": f"≤ {fmt(a)} °C"}
    return {"tmin": a - 5, "tmax": a + 5, "label": f"{fmt(a)} °C"}

AB_RE = re.compile(r"(acid|base)\s*:\s*(F\d{3})\s*,?\s*(?:ΔpKa|dpKa|pKa)?\s*(?:≥|>=)\s*([-−]?\d+(?:\.\d+)?)", re.I)
def acid_base(text):
    """'base: F019, ΔpKa ≥ -5' -> {"role": "base", "frag": "F019", "min": -5}"""
    m = AB_RE.search(str(text or ""))
    return {"role": m.group(1).lower(), "frag": m.group(2), "min": float(m.group(3).replace("−", "-"))} if m else None

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None

def cation_spec(text):
    """'[CX4][OH] | cation=1 | leaving=2 | then=eliminate' -> {"smarts", "cation", "leaving", "then", "add"}"""
    parts = [p.strip() for p in str(text or "").split("|") if p.strip()]
    if not parts:
        return None
    spec = {"smarts": parts[0], "cation": 1, "leaving": None, "then": "", "add": None}
    for p in parts[1:]:
        k, _, v = p.partition("=")
        k, v = k.strip(), v.strip()
        if k in ("cation", "leaving"):
            spec[k] = int(v)
        elif k == "then":
            spec["then"], _, add = v.partition(" ")
            spec["add"] = add.strip() or None
    return spec

rules = []
for r in rows("reaction_rules.xlsx", "Reactions"):
    keys = lambda col: KEY.findall(str(find(r, col)))
    rules.append({
        "id": r["Reaction ID"], "name": find(r, "Name"), "smarts": find(r, "Reaction SMARTS"),
        "reagents": keys("Reagent"), "byproducts": keys("By-product"), "conditions": find(r, "Conditions"),
        "consumes": frag_items(find(r, "Fragments consumed")), "forms": frag_items(find(r, "Fragment formed")),
        "acidBase": acid_base(find(r, "Acid/base needed")),
        "group": str(find(r, "Chemoselectivity group") or "").strip(),
        "cation": cation_spec(find(r, "Carbocation")),
        "code": str(find(r, "Reaction code") or ""),
        "equiv": (lambda v: "catalytic" if str(v).strip().lower().startswith("cat") else (num(v) if num(v) is not None else None))(find(r, "Reagent equivalents")),
        "stereo": {"outcome": str(find(r, "Stereo outcome") or ""), "cis": str(find(r, "Stereo template (cis") or ""),
                   "trans": str(find(r, "Stereo template (trans") or ""), "chiral": str(find(r, "Stereo template (chiral") or "")},
        "eas": ({"fc": "friedel" in str(find(r, "Aromatic substitution")).lower()} if str(find(r, "Aromatic substitution") or "").lower().startswith("yes") else None),
        # "a, b" = both needed; "a | b" = either is enough  ->  a list of alternatives, each a list of requirements
        "required": [alt for alt in ([requirement(c) for c in part.split(",") if c.strip()]
                                     for part in str(find(r, "Required conditions")).split("|")) if alt],
    })

# the electron flow of each rule (sheet "Mechanism"): one entry per curved-arrow step, with pKa-based proton transfers
def num(v):
    try: return float(v)
    except (TypeError, ValueError): return None
by_id = {rule["id"]: rule for rule in rules}
for rule in rules:
    rule["mechanism"] = []
for s in rows("reaction_rules.xlsx", "Mechanism"):
    if s["Reaction ID"] not in by_id or not isinstance(s.get("Step"), (int, float)):
        continue                                           # (the source note under the table)
    dpka, fpka = num(find(s, "Donor pKa")), num(find(s, "Acid formed pKa"))
    d = fpka - dpka if dpka is not None and fpka is not None else None
    by_id[s["Reaction ID"]]["mechanism"].append({
        "step": int(s["Step"]), "type": find(s, "Step type"), "from": find(s, "Electrons from"), "to": find(s, "Electrons to"),
        "made": find(s, "Bond made"), "broken": find(s, "Bond broken"),
        "donor": find(s, "Proton donor"), "donorPka": dpka, "formed": s.get("Acid formed") or "", "formedPka": fpka,
        "dpka": d, "K": 10 ** d if d is not None else None,
        "lg": s.get("Leaving group") or "", "lgPka": num(find(s, "Leaving group: pKa")),
        "smiles": find(s, "Intermediate"), "note": find(s, "Notes"),
        "before": find(s, "Structure the arrows act on"), "arrows": find(s, "Arrows ("),
    })
for rule in rules:
    rule["mechanism"].sort(key=lambda x: x["step"])
# a mechanism is optional: the product only needs the Reactions row; without Mechanism rows Advanced mode says so.
# Every mechanism is replayed (tools/check_mechanisms.py): its arrows must turn each step's start into its result.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from check_mechanisms import statuses
mech_status, mech_report = statuses()
print(mech_report)
for rule in rules:
    rule["mechStatus"] = (mech_status or {}).get(rule["id"], {}).get("status", "unchecked" if mech_status is None else "none")

# the named conditions (sheet "Conditions"): accepted words, the temperature range each stands for, what it counts as
words_of = lambda v: [w.strip().lower() for w in re.split(r",(?!\d)", str(v or "")) if w.strip()]   # 1,4-dioxane stays one word
conditions = {}
for c in rows("reaction_rules.xlsx", "Conditions"):
    conditions[str(c["Condition"]).strip().lower()] = {
        "words": words_of(c["Accepted words in the conditions box"]),
        "minC": num(c.get("Temperature from (°C)")), "maxC": num(c.get("Temperature to (°C)")),
        "implies": words_of(c.get("Counts as")),
    }
# the solvents (sheet "Solvents"): accepted words and boiling point ("reflux in <solvent>" runs at its boiling point)
solvents = {}
for s in rows("reaction_rules.xlsx", "Solvents"):
    if num(s.get("Boiling point (°C)")) is None:
        continue                                           # (the source note under the table)
    solvents[str(s["Solvent"]).strip()] = {"words": words_of(s["Accepted words in the conditions box"]),
                                           "smiles": s.get("SMILES") or "", "bp": num(s["Boiling point (°C)"]),
                                           "type": str(find(s, "Type") or "").strip()}
solvent_of = {k.lower(): k for k in solvents}
for rule in rules:
    for alt in rule["required"]:
        for q in alt:
            if "cond" in q and q["cond"] not in conditions and q["cond"] in solvent_of:
                q["solvent"] = solvent_of[q.pop("cond")]
unknown = sorted({q["cond"] for rule in rules for alt in rule["required"] for q in alt if "cond" in q and q["cond"] not in conditions}
                 | {i for c in conditions.values() for i in c["implies"] if i not in conditions})
if unknown:
    raise SystemExit("Required conditions that are not in the 'Conditions' or 'Solvents' sheet (and not a temperature such as 60-120 °C): " + ", ".join(unknown))
# conditions that come from a fragment (sheet "Fragment_Conditions"): checked for every fragment present
def cond_item(text):
    q = requirement(text)
    if "cond" in q and q["cond"] not in conditions and q["cond"] in solvent_of:
        q["solvent"] = solvent_of[q.pop("cond")]
    return q
fragment_conditions = []
for r in rows("reaction_rules.xlsx", "Fragment_Conditions"):
    fid = str(r["Fragment ID"]).strip()
    if not FID.fullmatch(fid):
        continue
    fragment_conditions.append({
        "frag": fid,
        "requires": [alt for alt in ([cond_item(c) for c in part.split(",") if c.strip()]
                                     for part in str(find(r, "Requires") or "").split("|")) if alt],
        "forbids": [cond_item(c) for c in str(find(r, "Forbids") or "").split(",") if c.strip()],
        "forbidFrags": FID.findall(str(find(r, "Forbidden fragments") or "")),
        "except": re.findall(r"R\d{3}", str(find(r, "Except") or "")),
        "reason": find(r, "Reason"),
    })
bad_fc = sorted({q["cond"] for fc in fragment_conditions for q in [x for alt in fc["requires"] for x in alt] + fc["forbids"]
                 if "cond" in q and q["cond"] not in conditions}
                | {f for fc in fragment_conditions for f in [fc["frag"]] + fc["forbidFrags"] if f not in fragments})
if bad_fc:
    raise SystemExit("Fragment_Conditions uses names that are not in the Conditions/Solvents/Fragments sheets: " + ", ".join(bad_fc))
# acids and bases (sheet "Acids_Bases"): any of these can serve a rule's "Acid/base needed" if its pKa is right
acids_bases = []
for r in rows("reaction_rules.xlsx", "Acids_Bases"):
    if r.get("Role") in ("acid", "base") and num(find(r, "pKa (")) is not None:
        acids_bases.append({"name": r["Name"], "smiles": find(r, "SMILES"), "role": r["Role"], "pKa": num(find(r, "pKa (")),
                            "conj": find(r, "Conjugate"), "note": find(r, "Notes")})
for rule in rules:
    ab = rule["acidBase"]
    if ab and ab["frag"] in fragments and fragments[ab["frag"]]["pKa" if ab["role"] == "base" else "pKaH"] is None:
        raise SystemExit(f"{rule['id']} needs the {'pKa' if ab['role'] == 'base' else 'conjugate-acid pKa'} of {ab['frag']}, which the Fragments sheet does not give")
# SN1/SN2/E1/E2 for alkyl halides: reagent classes (sheet "Nucleophiles") and the decision table ("Halide_Pathways")
nucleophiles = [{"name": r["Name"], "smiles": find(r, "SMILES"), "class": find(r, "Class"), "atom": find(r, "Attacking atom")}
                for r in rows("reaction_rules.xlsx", "Nucleophiles") if find(r, "Class")]
halide_pathways = []
for r in rows("reaction_rules.xlsx", "Halide_Pathways"):
    if r.get("Substrate (fragment class)") not in ("methyl", "1°", "2°", "3°"):
        continue
    halide_pathways.append({"substrate": r["Substrate (fragment class)"], "reagent": find(r, "Reagent class"), "solvent": find(r, "Solvent type"),
                            "heat": find(r, "Heat"), "major": find(r, "Major"), "minor": find(r, "Minor"), "why": find(r, "Why")})
classes = {n["class"] for n in nucleophiles} | {"any"}
badhp = sorted({h["reagent"] for h in halide_pathways if h["reagent"] not in classes}
               | {h["major"] for h in halide_pathways if h["major"] not in ("SN1", "SN2", "E1", "E2", "none")})
if badhp:
    raise SystemExit("Halide_Pathways uses reagent classes or pathways that are not known: " + ", ".join(badhp))
# condition suggestions for the conditions list (sheet "Condition_Suggestions")
condition_suggestions = [{"frag": str(r["Fragment ID"]).strip(), "suggest": find(r, "Suggest"), "why": find(r, "Why")}
                         for r in rows("reaction_rules.xlsx", "Condition_Suggestions") if FID.fullmatch(str(r["Fragment ID"]).strip()) and find(r, "Suggest")]
badcs = sorted({c["frag"] for c in condition_suggestions if c["frag"] not in fragments})
if badcs:
    raise SystemExit("Condition_Suggestions uses fragment IDs that are not on the Fragments sheet: " + ", ".join(badcs))
# chemoselectivity: fragments ranked within each group (sheet "Reactivity_Order")
reactivity = {}
for r in rows("reaction_rules.xlsx", "Reactivity_Order"):
    fid = str(r.get("Fragment ID") or "").strip()
    if not FID.fullmatch(fid) or num(find(r, "Rank")) is None:
        continue
    reactivity.setdefault(str(r["Chemoselectivity group"]).strip(), {})[fid] = {"rank": num(find(r, "Rank")), "why": find(r, "Why")}
badro = sorted({f for g in reactivity.values() for f in g if f not in fragments} |
               {rule["group"] for rule in rules if rule["group"] and rule["group"] not in reactivity})
if badro:
    raise SystemExit("Reactivity_Order / Chemoselectivity group problem (unknown fragment or group with no ranking): " + ", ".join(badro))
# directing groups for aromatic substitution (sheet "Directing_Groups")
WEIGHT = {"strong activator": 1000, "moderate activator": 300, "weak activator": 100, "weak deactivator": 30, "moderate deactivator": 10, "strong deactivator": 5}
directing = []
for r in rows("reaction_rules.xlsx", "Directing_Groups"):
    act = str(find(r, "Activation") or "").strip()
    if act not in WEIGHT:
        continue
    directing.append({"name": r["Group"], "smarts": find(r, "SMARTS"), "activation": act, "weight": WEIGHT[act],
                      "directs": str(find(r, "Directs")).strip(), "blocksFC": str(find(r, "Blocks Friedel") or "").strip().lower() == "yes",
                      "why": find(r, "Why")})
# known non-reactions (sheet "Known_No_Reaction")
known_none = [{"reagent": str(find(r, "Reagent (SMILES") or "").strip(), "name": find(r, "Reagent name"), "frag": str(r.get("Fragment ID") or "").strip(),
               "why": find(r, "Why")} for r in rows("reaction_rules.xlsx", "Known_No_Reaction") if FID.fullmatch(str(r.get("Fragment ID") or "").strip())]
badkn = sorted({k["frag"] for k in known_none if k["frag"] not in fragments})
if badkn:
    raise SystemExit("Known_No_Reaction uses fragment IDs that are not on the Fragments sheet: " + ", ".join(badkn))
# carbocation stability scores and shift priorities (sheet "Rearrangements")
rearrangements = {"stability": {}, "bonus": {}, "priority": {}}
for r in rows("reaction_rules.xlsx", "Rearrangements"):
    if r.get("Kind") in rearrangements and num(r.get("Value")) is not None:
        rearrangements[r["Kind"]][str(r["Item"])] = num(r["Value"])
nofrag = sorted({f for rule in rules for it in rule["consumes"] + rule["forms"] for f in it if f not in fragments}
                | {rule["acidBase"]["frag"] for rule in rules if rule["acidBase"] and rule["acidBase"]["frag"] not in fragments})
if nofrag:
    raise SystemExit("Fragment IDs used in reaction_rules.xlsx but missing from the Fragments sheet: " + ", ".join(nofrag))

missing = sorted({k for rule in rules for k in rule["reagents"] + rule["byproducts"] if k not in molecules})
if missing:
    raise SystemExit("These InChIKeys are used in reaction_rules.xlsx but missing from molecule_data.xlsx: " + ", ".join(missing))

# reagents (sheet "Reagents", learned from the recorded reactions): name, abbreviation and role of each reagent part
def _sheet_or_empty(book, name):
    try: return rows(book, name)
    except KeyError: return []
from rdkit import Chem as _Chem
def _canon(smi):
    m = _Chem.MolFromSmiles(str(smi or "")) if smi else None
    return _Chem.MolToSmiles(m) if m else ""
reagents = {}
for r in _sheet_or_empty("reaction_rules.xlsx", "Reagents"):
    k = _canon(r.get("SMILES (one reagent part)"))
    if k: reagents[k] = {"name": str(r.get("Name") or ""), "abbr": str(r.get("Abbreviation") or ""), "role": str(r.get("Role") or ""),
                         "also": str(r.get("Also") or ""), "pKaH": num(r.get("pKa of conjugate acid (bases)")),
                         "typed": str(r.get("SMILES (one reagent part)") or "").strip()}        # the part as the sheet types it: the Interferences advisories compare reagent parts as typed
# by-products (sheet "Byproducts"): what a learned reaction's leaving pieces really leave as
byproducts = []
for r in _sheet_or_empty("reaction_rules.xlsx", "Byproducts"):
    if not r.get("Leaving pieces (as computed)"): continue
    byproducts.append({"pieces": str(r["Leaving pieces (as computed)"]).strip(),
                       "when": ".".join(sorted(_canon(x) for x in str(r.get("When the reagent contains") or "").split(".") if x.strip())),
                       # "2*O" = two waters: the products box gets each molecule once, the text keeps the amounts
                       "byproducts": list(dict.fromkeys(_canon(re.sub(r"^\d+\*", "", x.strip())) for x in str(r.get("By-products (SMILES)") or "").split(".") if x.strip())),
                       "text": " + ".join(re.sub(r"^(\d+)\*(.+)$", r"\1 \2", x.strip()) for x in str(r.get("By-products (SMILES)") or "").split(".") if x.strip()),
                       "why": str(r.get("Why") or "")})
# interferences (sheet "Interferences"): what hinders or helps a template. A SITE row has a SMARTS whose first atom is the atom the template changes (and, for the ring
# substituents, a prefilter: SMARTS the product must contain somewhere); an ADVISORY row has a reagent side (alternatives "a | b", all of "x + y"; tokens frag:/smiles:/role:/solvent:)
# and group fragment IDs (js/interference.js)
interferences = []
_seen_int = set()
_TOKEN = re.compile(r"^(not:)?(role|frag|smiles|solvent|elem):\S.*$")
for r in _sheet_or_empty("reaction_rules.xlsx", "Interferences"):
    key = str(r.get("Key") or "").strip()
    if not key: continue
    if key in _seen_int: raise SystemExit("Interferences: the Key " + key + " is used twice")
    _seen_int.add(key)
    smarts = str(find(r, "Pattern") or "").strip()
    prefilter = str(find(r, "Prefilter") or "").strip()
    reagent_text = str(find(r, "Reagent side") or "").strip()
    reagent_alts = [[t.strip() for t in alt.split(" + ") if t.strip()] for alt in reagent_text.split(" | ") if alt.strip()] if reagent_text else []
    groups = FID.findall(str(find(r, "Groups (") or ""))
    if smarts:
        if not smarts.startswith("["): raise SystemExit("Interferences " + key + ": the Pattern must start with the bracket atom that reacts: " + smarts)
        if _Chem.MolFromSmarts(smarts) is None or _Chem.MolFromSmarts("[900" + smarts[1:]) is None: raise SystemExit("Interferences " + key + ": RDKit cannot read the Pattern " + smarts)
        if prefilter and _Chem.MolFromSmarts(prefilter) is None: raise SystemExit("Interferences " + key + ": RDKit cannot read the Prefilter " + prefilter)
    for alt in reagent_alts:
        for t in alt:
            if not _TOKEN.match(t): raise SystemExit("Interferences " + key + ": the reagent-side item '" + t + "' is not frag:/smiles:/role:/solvent:/elem:")
            tt = t[4:] if t.startswith("not:") else t
            if tt.startswith("frag:") and tt[5:] not in fragments: raise SystemExit("Interferences " + key + ": unknown fragment " + t)
    bad_g = [g for g in groups if g not in fragments]
    if bad_g: raise SystemExit("Interferences " + key + ": unknown fragment IDs " + ", ".join(bad_g))
    kind = str(find(r, "Kind") or "")
    if not smarts and not groups and kind != "data support" and str(find(r, "Needs") or "").strip() == "":
        raise SystemExit("Interferences " + key + ": needs a Pattern (a site row) or groups (an advisory)")
    interferences.append({"id": str(r.get("ID") or ""), "key": key, "name": str(find(r, "Name") or ""), "group": str(find(r, "Group") or ""), "kind": kind,
                          "anchor": str(find(r, "Applies to") or "") if smarts else "", "smarts": smarts, "prefilter": prefilter, "reagent": reagent_alts, "groups": groups,
                          "groupWhere": str(find(r, "Groups are looked for in") or "substrate").strip() or "substrate", "show": str(find(r, "Shown for") or "bystander").strip() or "bystander",
                          "minDistinct": int(num(find(r, "Distinct groups")) or 1), "needs": str(find(r, "Needs") or "").strip(), "page": str(find(r, "Used by the page") or "yes").strip().lower() != "no",
                          "effect_kind": str(find(r, "Effect kind") or ""), "where": str(find(r, "Where / when") or ""), "effect": str(find(r, "Effect on") or ""), "why": str(find(r, "Why") or ""),
                          "sigma": num(find(r, "Hammett")), "source": str(find(r, "Source") or "")})
out = os.path.join(ROOT, "data", "reaction_rules.js")
with open(out, "w", encoding="utf-8") as f:
    f.write("/* generated by tools/rules_to_js.py from reaction_rules.xlsx and molecule_data.xlsx: do not edit by hand */\n")
    f.write("window.REACTION_RULES = " + json.dumps({"rules": rules, "molecules": molecules, "conditions": conditions, "solvents": solvents, "fragments": fragments, "fragmentConditions": fragment_conditions, "acidsBases": acids_bases,
                                         "nucleophiles": nucleophiles, "halidePathways": halide_pathways,
                                         "conditionSuggestions": condition_suggestions, "reactivity": reactivity,
                                         "directing": directing, "knownNone": known_none, "rearrangements": rearrangements,
                                         "reagents": reagents, "byproducts": byproducts, "interferences": interferences}, ensure_ascii=False, indent=1) + ";\n")
print(f"wrote {out}: {len(rules)} rule(s), {len(molecules)} molecule(s), {len(fragments)} fragment(s)")
