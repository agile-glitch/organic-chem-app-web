"""Write the Fragment_Properties and Property_Guide sheets of molecule_data.xlsx from the generated data.

    python tools/fragment_properties_sheet.py

Reads the Fragments sheet, tools/data/fragment_properties.json (python tools/fragment_properties.py) and
tools/data/fragment_conditions.json (python tools/fragment_conditions.py). Writes nothing if a fragment lacks a row.
Fragment_Properties: one row per fragment, the properties that decide whether, where and how easily it reacts, in groups.
Property_Guide: one row per property: what it is, what it says about reactivity, where the value comes from, how to read it.
Both sheets are GENERATED: edit the generators or the Fragments sheet, not the cells.
"""
import json, os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
sys.path.insert(0, HERE)
BOOK = os.path.join(APP, 'molecule_data.xlsx')
PROPS = os.path.join(HERE, 'data', 'fragment_properties.json')
CONDS = os.path.join(HERE, 'data', 'fragment_conditions.json')

# (source, key, header, unit, group)   source: 'F' Fragments sheet, 'P' fragment_properties.json, 'C' fragment_conditions.json
COLUMNS = [
    ('F', 'Fragment ID', 'Fragment ID', '', 'Identity'),
    ('F', 'Fragment name', 'Fragment name', '', 'Identity'),
    ('F', 'Role', 'Role (Fragments sheet)', '', 'Identity'),
    ('F', 'Reactive atom(s)', 'Reactive atom(s) (Fragments sheet)', '', 'Identity'),
    ('P', 'model_smiles', 'Model compound (SMILES)', '', 'Model compound'),
    ('P', 'model_source', 'Where the model compound comes from', '', 'Model compound'),
    ('P', 'reacting_atom', 'Reacting atom in the model (element + atom number)', '', 'Model compound'),
    ('P', 'element', 'Electronic: element of the reacting atom', '', 'Electronic'),
    ('P', 'hybridisation', 'Electronic: hybridisation of the reacting atom', '', 'Electronic'),
    ('P', 'EN', 'Electronic: electronegativity of the reacting atom (Allred–Pauling)', 'χ', 'Electronic'),
    ('P', 'EN_neighbours', 'Electronic: mean electronegativity of its neighbours', 'χ', 'Electronic'),
    ('P', 'bond_polarity', 'Electronic: most polar bond at the reacting atom, Δχ (partner − atom; + means the atom is δ+)', 'χ', 'Electronic'),
    ('P', 'bond_polarity_partner', 'Electronic: partner atom of that bond', '', 'Electronic'),
    ('P', 'charge_gasteiger', 'Electronic: partial charge on the reacting atom, Gasteiger (the value the app uses per molecule)', 'e', 'Electronic'),
    ('P', 'charge_lowdin', 'Electronic: partial charge on the reacting atom, DFT Löwdin', 'e', 'Electronic'),
    ('P', 'oxidation_state', 'Electronic: oxidation state of the reacting atom', '', 'Electronic'),
    ('P', 'sigma_m', 'Electronic: Hammett σm of the fragment as a ring substituent', '', 'Electronic'),
    ('P', 'sigma_p', 'Electronic: Hammett σp', '', 'Electronic'),
    ('P', 'sigma_p_minus', 'Electronic: Hammett σp⁻ (through-conjugation)', '', 'Electronic'),
    ('P', 'sigma_character', 'Electronic: inductive / resonance character (from σm and σp)', '', 'Electronic'),
    ('P', 'E_homo_eV', 'Frontier orbitals: HOMO energy', 'eV', 'Frontier orbitals'),
    ('P', 'E_lumo_eV', 'Frontier orbitals: LUMO energy', 'eV', 'Frontier orbitals'),
    ('P', 'gap_eV', 'Frontier orbitals: HOMO–LUMO gap', 'eV', 'Frontier orbitals'),
    ('P', 'homo_share_pct', 'Frontier orbitals: share of the HOMO on the reacting atom (≈ f⁻)', '%', 'Frontier orbitals'),
    ('P', 'lumo_share_pct', 'Frontier orbitals: share of the LUMO on the reacting atom (≈ f⁺)', '%', 'Frontier orbitals'),
    ('P', 'homo_top', 'Frontier orbitals: atom carrying most of the HOMO', '', 'Frontier orbitals'),
    ('P', 'lumo_top', 'Frontier orbitals: atom carrying most of the LUMO', '', 'Frontier orbitals'),
    ('P', 'frontier_character', 'Frontier orbitals: character of the reacting atom (derived)', '', 'Frontier orbitals'),
    ('P', 'hardness_eV', 'Frontier orbitals: hardness η', 'eV', 'Frontier orbitals'),
    ('P', 'softness_per_eV', 'Frontier orbitals: softness S = 1/η', '1/eV', 'Frontier orbitals'),
    ('P', 'electrophilicity_eV', 'Frontier orbitals: electrophilicity index ω', 'eV', 'Frontier orbitals'),
    ('P', 'local_softness_minus', 'Frontier orbitals: local softness s⁻ at the reacting atom (nucleophilic)', '1/eV', 'Frontier orbitals'),
    ('P', 'local_softness_plus', 'Frontier orbitals: local softness s⁺ at the reacting atom (electrophilic)', '1/eV', 'Frontier orbitals'),
    ('P', 'polarisability_A3', 'Frontier orbitals: molecular polarisability α of the model compound', 'Å³', 'Frontier orbitals'),
    ('P', 'pKa', 'Acid–base: pKa of its most acidic H', '', 'Acid–base'),
    ('P', 'pKaH', 'Acid–base: pKa of its conjugate acid (how basic)', '', 'Acid–base'),
    ('P', 'weakest_base_that_deprotonates', 'Acid–base: weakest listed base that deprotonates it (ΔpKa ≥ 0)', '', 'Acid–base'),
    ('P', 'weakest_acid_that_protonates', 'Acid–base: weakest listed acid that protonates it (ΔpKa ≥ 0)', '', 'Acid–base'),
    ('P', 'steric_class', 'Steric: hindrance class at the reacting atom (the app\'s rule)', '', 'Steric'),
    ('P', 'steric_score', 'Steric: hindrance score (the app\'s rule)', '', 'Steric'),
    ('P', 'carbon_neighbours', 'Steric: carbon neighbours on the reacting atom', '', 'Steric'),
    ('P', 'substitution', 'Steric: degree of substitution (sp³ carbon centres)', '', 'Steric'),
    ('P', 'taft_Es_min', 'Steric: Taft Es of the bulkiest group on the reacting atom (model compound)', '', 'Steric'),
    ('P', 'taft_sigma_star_sum', 'Steric: Taft Σσ* of the groups on the reacting atom (model compound)', '', 'Steric'),
    ('P', 'buried_volume_pct', 'Steric: buried volume %V_bur (3.5 Å sphere round the reacting atom)', '%', 'Steric'),
    ('P', 'ortho_substituents', 'Steric: ortho substituents (aromatic reacting atom, model compound)', '', 'Steric'),
    ('F', 'Ring strain', 'Steric: ring strain (Fragments sheet)', '', 'Steric'),
    ('F', 'Leaving group', 'Leaving group: what leaves and how well (Fragments sheet)', '', 'Leaving group and bonds'),
    ('F', 'Made into a good leaving group by', 'Leaving group: how it is activated (Fragments sheet)', '', 'Leaving group and bonds'),
    ('P', 'bonds_text', 'Bonds: bonds at the reacting atom (average enthalpy; typical length)', '', 'Leaving group and bonds'),
    ('P', 'weakest_bond', 'Bonds: weakest bond at the reacting atom', 'kJ/mol', 'Leaving group and bonds'),
    ('P', 'weakest_bond_length', 'Bonds: typical length of that bond', 'Å', 'Leaving group and bonds'),
    ('C', 'n_consumed', 'Recorded conditions: recorded reactions that consumed it', 'n', 'Recorded conditions'),
    ('C', 'median_yield', 'Recorded conditions: median recorded yield', '%', 'Recorded conditions'),
    ('C', 'solvents', 'Recorded conditions: commonest solvents (share of those reactions)', '', 'Recorded conditions'),
    ('C', 'solvent_type', 'Recorded conditions: protic / polar aprotic / nonpolar split', '', 'Recorded conditions'),
    ('C', 'bases', 'Recorded conditions: commonest bases', '', 'Recorded conditions'),
    ('C', 'acids', 'Recorded conditions: commonest acids and Lewis acids', '', 'Recorded conditions'),
    ('C', 'catalysts', 'Recorded conditions: commonest metal catalysts', '', 'Recorded conditions'),
    ('C', 'other', 'Recorded conditions: other common reagents', '', 'Recorded conditions'),
    ('C', 'reacts_before', 'Competing sites: reacts in preference to (same molecule)', '', 'Competing sites'),
    ('C', 'gives_way_to', 'Competing sites: gives way to (same molecule)', '', 'Competing sites'),
]
FILLS = {'Identity': 'D9D9D9', 'Model compound': 'EDEDED', 'Electronic': 'DDEBF7', 'Frontier orbitals': 'E4DFEC', 'Acid–base': 'E2EFDA',
         'Steric': 'FCE4D6', 'Leaving group and bonds': 'F8CBAD', 'Recorded conditions': 'D0ECEC', 'Competing sites': 'FFF2CC'}

DFT = 'B3LYP/def2-SVP with a water continuum (ddCOSMO) on a GFN2-xTB geometry of the model compound'
# (group, property, what it is, what it says about reactivity, where the value comes from, limits)
GUIDE = [
    ('Model compound', 'Model compound', 'The simplest real molecule that contains the fragment, chosen by rule (tools/fragment_properties.py), not by hand.',
     'Every number below that is computed describes this molecule. The real value in a real molecule moves with its substituents.',
     'From the Molecules sheet, tools/data/reagents.json and the common small molecules of the recorded reactions: a molecule that contains the fragment plus at least one more heavy atom, with no stray charge (unless the fragment is an ion), no 3- or 4-membered ring, then as few extra rings, unsaturations and heteroatoms as possible, no formyl H on a carbonyl-type carbon, and the fewest heavy atoms.',
     'Read it with the model compound in view: ethyl vs methyl, an aryl vs an alkyl neighbour, change the numbers a little. tools/effects.py reads the same properties molecule by molecule.'),
    ('Electronic', 'Electronegativity of the reacting atom', 'Allred–Pauling electronegativity of the atom the app treats as reacting.',
     'With its neighbours\' values it sets bond polarity: an atom less electronegative than its neighbours is δ+ (electrophilic); a more electronegative one is δ− (nucleophilic, basic).',
     'js/mol_data.js (Allred 1961).', 'An atom property: ignores hybridisation (separate column) and what else is attached.'),
    ('Electronic', 'Mean electronegativity of its neighbours / most polar bond (Δχ)', 'The average electronegativity of the atoms bonded to the reacting atom, and the largest difference (partner − atom).',
     'A positive Δχ means the reacting atom is the electron-poor end of its most polar bond (a site for nucleophiles); a negative one means it is the electron-rich end.',
     'js/mol_data.js, applied to the bonds of the model compound (H atoms included).', 'Counts bonds, not resonance; a carbonyl carbon and an alkyl halide carbon both show a positive Δχ.'),
    ('Electronic', 'Partial charge (Gasteiger; DFT Löwdin)', 'Charge on the reacting atom: Gasteiger–Marsili (what the page shows per molecule) and a Löwdin population from the DFT density.',
     'More positive = better electrophilic site, more negative = better nucleophilic site, as a first guess. Charge does not decide the outcome alone (orbitals and sterics do).',
     'RDKit Gasteiger; Löwdin populations of the ' + DFT + '.', 'Atomic charges are method-dependent: compare within one column, not between the two. Gasteiger has no values for metals (blank).'),
    ('Electronic', 'Hybridisation', 'sp, sp², sp³ (or sp³d) of the reacting atom, from RDKit.',
     'More s character holds electrons tighter (alkyne C–H more acidic than alkene than alkane) and shortens and strengthens bonds; sp² centres are flat and open to attack from both faces.',
     'RDKit on the model compound.', 'Aromatic atoms are marked.'),
    ('Electronic', 'Oxidation state', 'Oxidation state of the reacting atom from electronegativity differences (bond electrons go to the more electronegative atom).',
     'Shows whether the centre can be oxidised or reduced and how far: a carbon at +3 (carboxylic acid) is already highly oxidised, at −3 (methyl) it can be oxidised in steps.',
     'Computed from the Allred–Pauling values of js/mol_data.js; aromatic bonds count 1.5, so a ".5" means delocalised.', 'A bookkeeping number, not a real charge.'),
    ('Electronic', 'Hammett σm, σp, σp⁻; inductive / resonance character', 'How the fragment, as a substituent on a benzene ring, changes the reactivity of the ring (σm mostly inductive; σp inductive + resonance; σp⁻ when it conjugates with a negative charge).',
     'Positive σ = electron-withdrawing (−I/−R): activates the ring to nucleophiles, deactivates it to electrophiles, acidifies a phenol or acid. Negative σ = donating. σp − σm says whether the resonance part pulls (≥ +0.05, −R) or pushes (≤ −0.10, +R).',
     'tools/pka.py SIGMA table (Hansch, Leo, Taft, Chem. Rev. 91, 165 (1991)), linked to the fragments that are that substituent.',
     'Blank where the app\'s table has no value for the group. The character label is a heuristic from σm and σp, not a measured σI / σR.'),
    ('Frontier orbitals', 'HOMO and LUMO energy; gap', 'Energy of the highest occupied and lowest unoccupied molecular orbital, and their difference.',
     'A high HOMO means the molecule gives electrons easily (nucleophile, base, oxidisable); a low LUMO means it accepts them easily (electrophile, reducible). A small gap means soft, polarisable and generally more reactive.',
     DFT + '. Kohn–Sham orbital energies.', 'Use them to RANK fragments, not as ionisation energies or electron affinities (B3LYP puts the HOMO several eV above the real IP). Anions sit high because of their charge.'),
    ('Frontier orbitals', 'HOMO / LUMO share on the reacting atom; atom carrying the most of each; character', 'How much of each frontier orbital sits on the reacting atom (Löwdin populations), where the orbital is largest, and a label from them.',
     'The coefficients say WHERE it reacts: the LUMO atom is attacked by nucleophiles (≈ Fukui f⁺), the HOMO atom attacks electrophiles (≈ Fukui f⁻). The label is electrophile-like when the LUMO share is ≥ 15% and ≥ 1.5× the HOMO share, nucleophile-like for the reverse, mixed when both are ≥ 15%, otherwise the orbitals sit elsewhere.',
     DFT + '. Shares are averaged over the orbitals within 0.05 eV of the HOMO (LUMO), so a degenerate pair (benzene, the p lone pairs of a halide, the π pairs of an alkyne) gives an answer that does not depend on how the pair is oriented.', 'Frontier-orbital approximation of the Fukui functions. If the largest atom is not the reacting atom (carboxylate: O, not C), the app\'s reacting-atom rule picked a different atom than the orbitals do. For saturated molecules the LUMO is often a diffuse σ*/Rydberg-like orbital (positive energy in this basis): read its share qualitatively.'),
    ('Frontier orbitals', 'Hardness η, softness S, local softness s⁻ / s⁺, electrophilicity ω', 'η = (E_LUMO − E_HOMO)/2; S = 1/η; local softness = S × orbital share; ω = μ²/2η with μ = (E_HOMO + E_LUMO)/2.',
     'Hard–hard and soft–soft pairs react best (HSAB): hard acids/bases are small, charged, low polarisability; soft ones are large and polarisable. s⁺ and s⁻ locate the soft electrophilic and nucleophilic site; ω ranks electrophilic strength.',
     'Computed from the frontier-orbital energies and shares (Parr, Pearson).', 'Derived from Kohn–Sham energies: relative scales only.'),
    ('Frontier orbitals', 'Polarisability α', 'How strongly an electric field distorts the electron cloud of the model compound (static, isotropic).',
     'More polarisable = softer: stabilises charge-transfer transition states and prefers soft partners (sp³ carbon, Pd). In protic solvents the soft, large ion is also the better nucleophile (I⁻ > Br⁻ > Cl⁻), because it is less tightly solvated; in polar aprotic solvents that order reverses.',
     'Finite electric field (forward difference, 0.002 a.u.) in the same DFT set-up.', 'A molecular property of the whole model compound, so it grows with its size; compare fragments with similarly sized models.'),
    ('Acid–base', 'pKa and pKa of the conjugate acid', 'Acidity of the most acidic H and basicity (pKaH) of the fragment, typical values in water.',
     'Decides proton transfer: a base removes the H if its conjugate acid is weaker than the fragment (ΔpKa = pKaH(base) − pKa(fragment) ≥ 0). A low pKaH means a poor nucleophile/base; protonation of C=O (pKaH about −6) is what acid catalysis uses.',
     'Fragments sheet (Evans pKa table; Clayden), read by tools/pka.py.', 'Typical values; the real pKa moves with substituents (tools/pka.py Hammett estimates for phenols, benzoic acids, anilines).'),
    ('Acid–base', 'ΔpKa: weakest listed base / acid that does it', 'ΔpKa is the pKa gap between the fragment and a reagent. It is not a property of the fragment alone, so it is stored as the weakest reagent of the Acids_Bases sheet that fully deprotonates (or protonates) the fragment, with its ΔpKa.',
     'Tells which reagents can act on the fragment at all by proton transfer, and how strongly (ΔpKa ≥ 0 goes to completion; −5 ≤ ΔpKa < 0 can still work as an equilibrium, as in the app\'s rules).',
     'Acids_Bases sheet of reaction_rules.xlsx and the pKa columns above.', 'Only the 25 reagents of that sheet; "none in the table" means a stronger base or acid than any listed is needed.'),
    ('Steric', 'Hindrance class and score; carbon neighbours; substitution', 'How crowded the reacting atom is: the app\'s rule (heavy neighbours + the most crowded neighbour\'s own branches), the number of carbon neighbours, and 1°/2°/3° for sp³ carbon.',
     'Crowding slows SN2 and addition (methyl > 1° > 2° > 3°), favours elimination and SN1 at 3° centres, and shields a site from large reagents.',
     'tools/effects.py (the same rule as the page) on the model compound.', 'Counts atoms, not 3-D shape; the model compound\'s own groups set the value.'),
    ('Steric', 'Taft Es and Σσ*', 'Taft steric (Es) and polar (σ*) parameters of the groups R on the reacting atom: Es more negative = bulkier; σ* > 0 pulls electrons, < 0 pushes.',
     'Es predicts how a bulky R slows attack at the centre; σ* predicts how R changes its electrophilicity or basicity.',
     'tools/effects.py TAFT table (Taft 1956; Hansch, Leo, Taft 1991) applied to the model compound.', 'Taft parameters belong to the substituent R, so they depend on the molecule: the value is that of the model compound\'s groups. Blank when the centre carries no tabulated group.'),
    ('Steric', 'Buried volume %V_bur', 'Percentage of a 3.5 Å sphere around the reacting atom that the other atoms of the molecule fill (van der Waals radii × 1.17).',
     'A direct 3-D measure of how shielded the site is: higher = harder for a reagent to reach. It stands in for cone angle, which is defined for ligands on a metal, not for organic centres.',
     'Computed on the GFN2-xTB geometry; radii from js/mol_data.js.', 'The reacting atom itself is left out; crowding comes from the model compound\'s substituents, so a larger real molecule is usually more buried.'),
    ('Steric', 'Ortho substituents', 'Number of substituents on the two ring atoms next to an aromatic reacting atom.',
     'Ortho groups block attack at the ipso carbon, twist substituents out of conjugation and slow cross-coupling and SNAr.',
     'Counted on the model compound (0 for the plain ring).', 'Blank unless the reacting atom is aromatic; the real value is per molecule (tools/effects.py counts ortho groups in its steric score). Peri crowding is not tracked.'),
    ('Steric', 'Ring strain', 'Strain energy of a ring fragment (epoxide about 27 kcal/mol).', 'Strained rings open easily with nucleophiles or acids; 5- and 6-membered rings behave like open chains.',
     'Fragments sheet (Anslyn & Dougherty).', 'Copied text; "none" for fragments that are not rings.'),
    ('Leaving group and bonds', 'Leaving group; how it is activated', 'What leaves from the fragment and how good a leaving group it is, and the reagent or step that turns a poor one into a good one.',
     'Leaving-group ability tracks the pKa of its conjugate acid (I⁻ > Br⁻ > Cl⁻ > F⁻; TfO⁻ > TsO⁻ > MsO⁻; HO⁻ and RO⁻ are poor until protonated or sulfonylated).',
     'Fragments sheet (copied).', 'Text only; the numeric pKa of each conjugate acid is in the Acids_Bases and pKa data, not here.'),
    ('Leaving group and bonds', 'Bonds at the reacting atom; weakest bond and its length', 'Every bond at the reacting atom with its average bond enthalpy and typical length, and the weakest of them.',
     'Weak bonds break first (average enthalpies in this table: C–I 240 < C–Br 275 < C–Cl 330 < C–F 439 kJ/mol); short, strong bonds (C=O 741, C≡N 891) resist; bond length goes with strength.',
     'Average bond enthalpies (OpenStax Chemistry 2e) and typical lengths (Allen et al. 1987) of js/mol_data.js.', 'Averages over many molecules, not bond dissociation energies of this molecule. Blank when a bond at the atom has no tabulated enthalpy (aromatic bonds, and C–Mg, C–Sn, Al–H and similar). Charge-separated drawings (nitro, N-oxide, carboxylate) use localised bond orders. The bond that breaks in a given reaction is not always the weakest.'),
    ('Recorded conditions', 'Reactions consumed, median yield, solvents, solvent type, bases, acids, catalysts, other reagents', 'What the 1.1 million recorded reactions (US patents, Open Reaction Database) did with the fragment: how many reactions consumed it, their median yield, the commonest solvents, bases, acids and Lewis acids, metal catalysts and other reagents, each as a share of those reactions.',
     'Shows the environment a fragment reacts in: protic vs polar aprotic solvent, base- or acid-mediated, metal-catalysed (for example Pd/C and Fe for nitro reduction; carbonate and Pd(PPh₃)₄ for aryl bromides).',
     'tools/fragment_conditions.py over data/learned/rx.', 'A count of what chemists recorded, with the biases of patent chemistry. TEMPERATURE, CONCENTRATION, TIME AND EQUIVALENTS ARE NOT IN THE SOURCE DATA, so they are not reported. A fragment consumed in a reaction may not be the reactive site the chemist intended.'),
    ('Competing sites', 'Reacts in preference to; gives way to', 'Other fragments in the same starting-material molecule that survived while this one was consumed (it reacts in preference to them), and the ones that were consumed while this one survived (it gives way to them).',
     'Chemoselectivity: which group reacts first when several are present. Each entry reads: partner [share of the reactions that consumed this fragment in which the partner survived of n vs share of the reactions that consumed the partner in which this fragment survived of n].',
     'tools/fragment_conditions.py, sampled up to 2,500 reactions per fragment; needs ≥ 15 reactions each side, a ≥ 60% survival and a ≥ 30-point difference (one-sided: ≥ 90% of ≥ 30).', 'The chemist picked reagents to favour one site, so this is chemoselectivity under the conditions people used, not an intrinsic rate. Fragments that are rarely consumed (pyridine N, arene C–H) appear often as "survivors".'),
    ('Not stored per fragment', 'Cone angle', 'Tolman cone angle: the width of the cone a ligand fills round a metal.', 'Buried volume is used instead (same purpose, defined for any centre).', '-', 'Not defined for organic reacting centres.'),
    ('Not stored per fragment', 'Temperature, concentration, time', 'Reaction conditions.', 'They decide rate and selectivity (kinetic vs thermodynamic control, competing pathways).', '-', 'Not recorded in the source data; they belong to a reaction, not a fragment. Solvent, base, acid and catalyst are reported above.'),
]


def sheet_rows(name):
    import openpyxl
    ws = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)[name]
    it = ws.iter_rows(values_only=True); head = [str(h or '').strip() for h in next(it)]
    return [dict(zip(head, r)) for r in it if r and r[0] not in (None, '')]


def build():
    frags = [f for f in sheet_rows('Fragments') if str(f['Fragment ID']).startswith('F')]
    props = {r['Fragment ID']: r for r in json.load(open(PROPS, encoding='utf-8'))}
    conds = json.load(open(CONDS, encoding='utf-8'))
    missing = [f['Fragment name'] for f in frags if f['Fragment ID'] not in props or props[f['Fragment ID']].get('error') or f['Fragment ID'] not in conds]
    if missing: raise SystemExit('no properties / conditions for: ' + ', '.join(missing) + '\nrun tools/fragment_properties.py and tools/fragment_conditions.py first')
    rows = []
    for f in frags:
        fid = f['Fragment ID']; src = {'F': f, 'P': props[fid], 'C': conds[fid]}
        rows.append([src[s].get(k) for s, k, _, _, _ in COLUMNS])
    return rows


def write():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    rows = build()
    wb = openpyxl.load_workbook(BOOK)
    for n in ('Fragment_Properties', 'Property_Guide'):
        if n in wb.sheetnames: del wb[n]
    at = wb.sheetnames.index('Fragments') + 1
    ws = wb.create_sheet('Fragment_Properties', at)
    ws.append([f'{h}' + (f' [{u}]' if u else '') for _, _, h, u, _ in COLUMNS])
    for r in rows: ws.append(r)
    wrap = Alignment(wrap_text=True, vertical='top')
    for j, (_, k, h, u, g) in enumerate(COLUMNS, 1):
        c = ws.cell(1, j); c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical='center'); c.fill = PatternFill('solid', fgColor=FILLS[g])
        long = k in ('Role', 'Reactive atom(s)', 'Leaving group', 'Made into a good leaving group by', 'bonds_text', 'solvents', 'solvent_type', 'bases', 'acids', 'catalysts', 'other',
                     'reacts_before', 'gives_way_to', 'weakest_base_that_deprotonates', 'weakest_acid_that_protonates', 'model_source', 'Fragment name')
        ws.column_dimensions[get_column_letter(j)].width = 46 if long else 24 if k in ('model_smiles', 'sigma_character', 'frontier_character', 'hybridisation') else 16
        for i in range(2, len(rows) + 2): ws.cell(i, j).alignment = wrap
    ws.row_dimensions[1].height = 78; ws.freeze_panes = 'C2'
    gs = wb.create_sheet('Property_Guide', at + 1)
    gs.append(['Group', 'Property (column)', 'What it is', 'What it says about reactivity', 'Where the value comes from', 'Limits: how to read it'])
    for r in GUIDE: gs.append(list(r))
    for j, w in enumerate([22, 34, 55, 70, 55, 65], 1):
        gs.column_dimensions[get_column_letter(j)].width = w
        gs.cell(1, j).font = Font(bold=True); gs.cell(1, j).fill = PatternFill('solid', fgColor='D9D9D9')
        for i in range(2, len(GUIDE) + 2): gs.cell(i, j).alignment = wrap
    gs.freeze_panes = 'C2'
    how = wb['How to use']
    marker = 'Fragment_Properties:'
    if not any(isinstance(r[0].value, str) and r[0].value.startswith(marker) for r in how.iter_rows()):
        how.append([marker + ' one row per fragment: electronic, frontier-orbital, acid-base, steric, leaving-group / bond and recorded-condition properties, computed (python tools/fragment_properties.py, tools/fragment_conditions.py) and written by tools/fragment_properties_sheet.py. Generated: do not edit the cells.'])
        how.append(['Property_Guide: one row per property: what it is, what it says about reactivity, where the value comes from, and its limits (temperature and concentration are not in the recorded data; cone angle is replaced by buried volume).'])
    wb.save(BOOK)
    print(f'wrote Fragment_Properties ({len(rows)} fragments x {len(COLUMNS)} columns) and Property_Guide ({len(GUIDE)} rows) to {BOOK}')


if __name__ == '__main__':
    write()
