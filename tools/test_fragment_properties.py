"""Test the per-fragment property data:   python tools/test_fragment_properties.py

1. every fragment has a model compound that really contains it, and a complete, finite row of properties
2. the numbers behave chemically (carbonyl LUMO sits on the carbonyl carbon, alkyl iodide LUMO on C-I, Grignard HOMO on carbon ...)
3. the Fragment_Properties and Property_Guide sheets of molecule_data.xlsx match the generated data, and the sheet's IDs match the Fragments sheet
4. after a re-ID ('fragment_id.py --fix') the generated rows of that fragment are forgotten, not renamed
"""
import json, math, os, shutil, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import fragment_id as FI, fragment_properties as FP, fragment_properties_sheet as FS
from rdkit import Chem
if not (os.path.exists(FS.PROPS) and os.path.exists(FS.CONDS)):
    print('SKIP: the property data is not generated yet (python tools/fragment_properties.py; python tools/fragment_conditions.py; '
          'python tools/fragment_properties_sheet.py)')
    sys.exit(0)
fails = []
def check(ok, what):
    print(("PASS  " if ok else "FAIL  ") + what)
    if not ok: fails.append(what)

frags = FP.fragments(); by_id = {f['Fragment ID']: f for f in frags}
props = {r['Fragment ID']: r for r in json.load(open(FS.PROPS, encoding='utf-8'))}
conds = json.load(open(FS.CONDS, encoding='utf-8'))
name = lambda n: next(i for i, f in by_id.items() if str(f['Fragment name']).startswith(n))

# 1 ---------------------------------------------------------------------------------------------------------------
check(set(props) == set(by_id), f'a properties row for each of the {len(by_id)} fragments')
check(not [r for r in props.values() if r.get('error')], 'no fragment failed')
bad = [by_id[i]['Fragment name'] for i, r in props.items() if not Chem.MolFromSmiles(r['model_smiles']).HasSubstructMatch(Chem.MolFromSmarts(by_id[i]['SMARTS pattern']))]
check(not bad, 'every model compound contains its fragment' + (f' (not: {bad})' if bad else ''))
need = ['EN', 'hybridisation', 'oxidation_state', 'E_homo_eV', 'E_lumo_eV', 'gap_eV', 'homo_share_pct', 'lumo_share_pct', 'hardness_eV', 'softness_per_eV',
        'electrophilicity_eV', 'polarisability_A3', 'charge_lowdin', 'buried_volume_pct', 'steric_class', 'bonds_text']
gaps = [(r['Fragment name'][:25], k) for r in props.values() for k in need if r.get(k) is None]
check(not gaps, 'the always-defined properties are all filled' + (f' (missing: {gaps[:5]})' if gaps else ''))
nums = [(r['Fragment name'][:25], k) for r in props.values() for k, v in r.items() if isinstance(v, float) and not math.isfinite(v)]
check(not nums, 'no NaN or infinite number')
check(all(r['E_homo_eV'] < r['E_lumo_eV'] for r in props.values()), 'HOMO below LUMO for every fragment')
check(all(0 <= r['homo_share_pct'] <= 100 and 0 <= r['lumo_share_pct'] <= 100 for r in props.values()), 'orbital shares between 0 and 100 %')
check(all(r['polarisability_A3'] > 0 and 0 < r['buried_volume_pct'] < 100 for r in props.values()), 'polarisability positive, buried volume between 0 and 100 %')
check(all(abs(sum(r['raw']['homo_share']) - 1) < 1e-3 and abs(sum(r['raw']['lumo_share']) - 1) < 1e-3 for r in props.values()), 'each orbital\'s atomic shares add up to 1')

# 2 ---------------------------------------------------------------------------------------------------------------
P = lambda n: props[name(n)]
for n in ('ketone', 'aldehyde'):
    r = P(n); check(r['lumo_share_pct'] > r['homo_share_pct'] and r['frontier_character'].startswith('electrophile'), f'{n}: the LUMO sits on the carbonyl carbon (electrophile-like): {r["lumo_share_pct"]}% vs HOMO {r["homo_share_pct"]}%')
r = P('alkyl iodide'); check(r['lumo_share_pct'] > 15 and r['lumo_top'].startswith(('C', 'I')) and r['homo_top'].startswith('I'), f'alkyl iodide: HOMO on I, LUMO (σ* C–I) on C/I ({r["homo_top"]}, {r["lumo_top"]})')
r = P('Grignard'); check(r['homo_top'].startswith('C') and r['lumo_top'].startswith('Mg'), f'Grignard: HOMO on carbon, LUMO on Mg ({r["homo_top"]}, {r["lumo_top"]})')
r = P('carboxylate anion'); check(r['homo_top'].startswith('O'), f'carboxylate: the HOMO is on oxygen ({r["homo_top"]})')
r = P('alcohol (1°)'); check(r['homo_share_pct'] > r['lumo_share_pct'] and r['frontier_character'].startswith('nucleophile'), f'alcohol O: HOMO-like (nucleophile): {r["frontier_character"]}')
check(P('carboxylic acid')['oxidation_state'] == 3.0 and P('alcohol (1°)')['oxidation_state'] == -2.0, 'oxidation states: acid C +3, alcohol O −2')
check(P('alkyl iodide')['weakest_bond'].startswith('C–I') and P('alkyl iodide')['weakest_bond_kJ'] < P('alkyl bromide')['weakest_bond_kJ'] < P('alkyl chloride')['weakest_bond_kJ'], 'weakest bond: C–I < C–Br < C–Cl')
check(P('nitroarene')['sigma_p'] > 0.7 and P('phenol')['sigma_p'] < 0 and P('nitroarene')['sigma_character'].startswith('−I'), 'Hammett: nitro withdraws, hydroxy donates')
check(P('alkyl iodide')['polarisability_A3'] > P('alkyl chloride')['polarisability_A3'], 'iodoethane is more polarisable than chloroethane')
check(P('ketone')['gap_eV'] < P('alcohol (1°)')['gap_eV'], 'ketone gap smaller than alcohol gap')
r = conds[name('nitroarene')]; check(r['n_consumed'] > 1000 and 'Pd' in (r['catalysts'] or ''), 'recorded nitro-arene reactions: Pd catalysts are among the commonest')
r = conds[name('Grignard')]; check('tetrahydrofuran' in r['solvents'] and r['solvent_type'].startswith('polar aprotic'), 'recorded Grignard reactions: THF, polar aprotic')

# 3 ---------------------------------------------------------------------------------------------------------------
import openpyxl
wb = openpyxl.load_workbook(FS.BOOK, read_only=True, data_only=True)
check({'Fragment_Properties', 'Property_Guide'} <= set(wb.sheetnames), 'the workbook has the Fragment_Properties and Property_Guide sheets')
if 'Fragment_Properties' in wb.sheetnames:
    rows = list(wb['Fragment_Properties'].iter_rows(values_only=True))
    check(len(rows[0]) == len(FS.COLUMNS) and len(rows) == len(frags) + 1, f'{len(frags)} fragment rows x {len(FS.COLUMNS)} columns')
    check([r[0] for r in rows[1:]] == [f['Fragment ID'] for f in frags], 'the sheet lists the Fragments-sheet IDs in order')
    j = [c[1] for c in FS.COLUMNS].index('E_homo_eV')
    check(all(abs(r[j] - props[r[0]]['E_homo_eV']) < 1e-9 for r in rows[1:]), 'sheet HOMO energies equal the generated data')
    g = list(wb['Property_Guide'].iter_rows(values_only=True))
    check(len(g) == len(FS.GUIDE) + 1 and all(all(c for c in r) for r in g[1:]), 'every guide row is complete')
check(not FI.stale(), 'every fragment ID still matches its SMARTS')

# 4 ---------------------------------------------------------------------------------------------------------------
tmp = tempfile.mkdtemp(); real = FI.ROOT
try:
    os.makedirs(os.path.join(tmp, 'tools', 'data'))
    for f in FI.DATA_BY_ID: shutil.copy(os.path.join(real, f), os.path.join(tmp, f))
    FI.ROOT = tmp; victim = name('ketone')
    FI.forget({victim})
    left = json.load(open(os.path.join(tmp, 'tools', 'data', 'fragment_properties.json'), encoding='utf-8'))
    check(victim not in {r['Fragment ID'] for r in left} and len(left) == len(frags) - 1, 're-ID: the properties row of the changed fragment is forgotten')
    check(victim not in json.load(open(os.path.join(tmp, 'tools', 'data', 'fragment_models.json'), encoding='utf-8')) and victim not in json.load(open(os.path.join(tmp, 'tools', 'data', 'fragment_conditions.json'), encoding='utf-8')), 're-ID: model and conditions rows forgotten too')
finally:
    FI.ROOT = real; shutil.rmtree(tmp, ignore_errors=True)

print("\n" + ("ALL TESTS PASSED" if not fails else f"{len(fails)} FAILED"))
sys.exit(1 if fails else 0)
