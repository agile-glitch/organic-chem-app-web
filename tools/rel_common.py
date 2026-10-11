"""Shared pieces of the template-relevance work (tools/rel_features.py, rel_train.py, eval_top1.py).

  rows       the learned reactions (learned_reactions.xlsx), one list each:
             [Learned ID, starting materials, reagent, product, template #, reagent key, reagent part keys, page can read,
              patent / ORD id, look-up key]
  the split  by PATENT (md5 of the patent id mod 20): 1 patent in 20 is the TEST set, another 1 in 20 the VALIDATION set, the rest
             is training. Every reaction of a held-out patent is kept out of everything that is learned (template counts, the
             relevance model, win rates, the neighbour search). A patent usually records many near-identical reactions, so
             splitting by reaction id (tools/learn_selectivity.py) lets a test reaction have a twin in training; splitting by
             patent does not. All accuracy figures of this work are on the test patents; fusion weights and calibration are
             fitted on the validation patents only.

The workbook is git-ignored (240 MB); in a git worktree it is looked for in the folders above (LEARNED_XLSX overrides).
The parsed rows are cached in the temp folder (organic_chem_eval_rows.pkl).
"""
import hashlib, json, os, pickle, tempfile, time

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
CACHE = os.path.join(tempfile.gettempdir(), 'organic_chem_eval_rows.pkl')
HELD_OUT_ONE_IN = 20


def workbook_path():
    p = os.environ.get('LEARNED_XLSX') or os.path.join(APP, 'learned_reactions.xlsx')
    if os.path.exists(p): return p
    up = APP
    for _ in range(5):                                   # a git worktree does not have the (git-ignored) workbook
        up = os.path.dirname(up)
        if os.path.exists(os.path.join(up, 'learned_reactions.xlsx')): return os.path.join(up, 'learned_reactions.xlsx')
    return p


def patent_of(source):
    s = str(source or '').strip()
    return s.split()[0] if s else '?'


def is_held_out_patent(patent):
    return int(hashlib.md5(str(patent).encode()).hexdigest(), 16) % HELD_OUT_ONE_IN == 0


TRAIN, VAL, TEST = 0, 1, 2


def split_of_patent(patent):
    """0 train, 1 validation (fusion weights, calibration are fitted here), 2 test (the reported figures; nothing is ever tuned on it).
    The test patents are the same 1 in 20 that the earlier evaluation of the learned tier (the 45.2 % baseline) held out."""
    h = int(hashlib.md5(str(patent).encode()).hexdigest(), 16) % HELD_OUT_ONE_IN
    return TEST if h == 0 else VAL if h == 1 else TRAIN


def split_codes(rows):
    cache, out = {}, []
    for r in rows:
        p = r[8]
        v = cache.get(p)
        if v is None: v = cache[p] = split_of_patent(p)
        out.append(v)
    return out


def read_rows():
    """{'rows': [...], 'smarts': [template SMARTS by template #], 'tid': [template IDs]}  (cached)"""
    xlsx = workbook_path()
    if os.path.exists(CACHE) and (not os.path.exists(xlsx) or os.path.getmtime(CACHE) > os.path.getmtime(xlsx)):
        return pickle.load(open(CACHE, 'rb'))
    import openpyxl
    t0 = time.time()
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    trows = [r for r in list(wb['Templates'].iter_rows(values_only=True))[1:] if r and r[0]]
    tindex = {r[0]: k for k, r in enumerate(trows)}
    rows = []
    for name in wb.sheetnames:
        if not name.startswith('Learned'): continue
        head = None
        for r in wb[name].iter_rows(values_only=True):
            if head is None: head = {h: k for k, h in enumerate(r)}; continue
            if not r or not r[0]: continue
            g = lambda col: r[head[col]]
            rows.append([g('Learned ID'), (g('Starting materials') or '').split('.'), g('Reagent') or '', g('Product'), tindex[g('Template ID')],
                         g('Reagent key (InChIKey of its distinct parts)') or '', (g('Reagent part keys') or '').split(), g('Page can read every molecule') == 'yes',
                         patent_of(g('Source (patent / ORD id)')), g('Look-up key (InChIKey of starting materials + reagent parts)') or ''])
        print(f'  read {name}: {len(rows):,} rows ({time.time() - t0:.0f} s)', flush=True)
    data = {'rows': rows, 'smarts': [r[1] for r in trows], 'tid': [r[0] for r in trows]}
    pickle.dump(data, open(CACHE, 'wb'), protocol=4)
    return data


def rows_subset(indices):
    """{reaction index: row} for these reactions only, cached in a small file keyed by the set (the full rows need 1.5 GB of memory; a ranker run needs 27,000 of them)"""
    idx = sorted(set(int(i) for i in indices))
    key = hashlib.md5(repr(idx).encode()).hexdigest()[:12]
    path = os.path.join(tempfile.gettempdir(), f'organic_chem_rows_subset_{key}.pkl')
    if os.path.exists(path): return pickle.load(open(path, 'rb'))
    rows = read_rows()['rows']
    out = {i: rows[i] for i in idx}
    pickle.dump(out, open(path, 'wb'), protocol=4)
    return out


def held_out_flags(rows):
    cache, out = {}, []
    for r in rows:
        p = r[8]
        v = cache.get(p)
        if v is None: v = cache[p] = is_held_out_patent(p)
        out.append(v)
    return out


def fnv1a32(text):
    h = 0x811C9DC5
    for ch in text.encode('utf-8'):
        h = ((h ^ ch) * 0x01000193) & 0xFFFFFFFF
    return h


_SOLVENT_LIKE = None


def solvent_like():
    """canonical SMILES of the solvents and common agents of data/names.js (js/precedent.js SOLVENT_LIKE): a record whose 'product' is one of
    them is a data error, not a miss"""
    global _SOLVENT_LIKE
    if _SOLVENT_LIKE is None:
        from rdkit import Chem
        s = open(os.path.join(APP, 'data', 'names.js'), encoding='utf-8').read()
        j = json.loads(s[s.index('=') + 1:s.rindex('}') + 1])
        out = set()
        for smi in list(j['solvents']) + list(j['agents']) + ['O', 'CO', 'CCO', 'CC(=O)O', 'CCOC(C)=O']:
            m = Chem.MolFromSmiles(smi)
            if m is not None: out.add(Chem.MolToSmiles(m))
        _SOLVENT_LIKE = out
    return _SOLVENT_LIKE


def gradable(row):
    """the Validation tab's data hygiene (js/precedent.js plausible): the recorded product is not a solvent or common agent, and not a fragment
    smaller than 40 % of the largest starting material (when that has 10+ heavy atoms). A secondary figure only: the headline counts every record."""
    from rdkit import Chem
    p = Chem.MolFromSmiles(row[3])
    if p is None: return False
    if Chem.MolToSmiles(p) in solvent_like(): return False
    big = max((m.GetNumHeavyAtoms() for m in (Chem.MolFromSmiles(x) for x in row[1]) if m is not None), default=0)
    return not (big >= 10 and p.GetNumHeavyAtoms() < 0.4 * big)
