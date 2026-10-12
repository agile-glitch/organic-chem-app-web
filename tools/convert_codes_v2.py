"""Rewrite every stored version 1 reaction code (starting materials~reagent~conditions~products) as version 2
(reactants>reagents>products | cond=...), the format js/rxn_code.js and tools/rxn_code.py define.

    python tools/convert_codes_v2.py            # convert, then: python tools/rules_to_js.py   (data/reaction_rules.js from the workbook)
    python tools/convert_codes_v2.py --check    # only count what is still version 1

What it rewrites:
  reaction_rules.xlsx         sheet Reactions, column "Reaction code" (and the sentence describing the code on sheet "How to use")
  tools/reaction_tests.json   the code of every workbook reaction and feature case (except the one case that checks an old code is still read)
  tools/learned_tests.json    the sample of learned reactions for "Run all tests"
  data/learned/rx/*.js        the 1.09 million recorded reactions (column 1 of each row)

Each code is rewritten from its own fields (same molecules, same conditions text, nothing re-canonicalised), and the new code must read
back to exactly the fields of the old one, or nothing is written. Running it again changes nothing.
"""
import base64, gzip, json, os, re, sys

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import rxn_code  # noqa: E402

KEEP_V1 = {'the old ~ code is read and rewritten as v2'}           # reaction_tests.json: the case that tests the version 1 reader
HOW_TO = ("Reaction code: identifies one specific reaction as 'reactants>reagents>products | conc=<strength> | cond=<conditions>'. "
          "The first word is an ordinary reaction SMILES that RDKit and other chemistry programs read as it stands; the molecule parts are "
          "RDKit canonical SMILES, in the SMILES-line order, so the same reaction always gives the same code. After a space come what a "
          "reaction SMILES cannot say: conc= the reagent's strength (37%w/w, 2M, conc, dil) and cond= the conditions box's comma list. "
          "The Reactions page shows the code in the bar at the top, and a code pasted there (Enter) fills the page; the old form "
          "'starting materials~reagent~conditions~products' is still read there. Each rule's code is the reaction written in its row.")


def convert(code, where):
    """version 2 for a version 1 code (a version 2 code or an empty one is returned unchanged); checked to read back the same"""
    if not code or rxn_code.parse(code).get('version') != 1: return code, False
    old = rxn_code.parse(code)
    for k in ('cond',):
        if '|' in old[k]: raise ValueError(f'{where}: the conditions contain "|", which version 2 cannot hold: {code!r}')
    new = rxn_code.to_v2(code)
    back = rxn_code.parse(new)
    same = all(back[k] == old[k] for k in ('reactants', 'agents', 'products')) and back['cond'] == old['cond'].strip() and back['version'] == 2
    if not same: raise ValueError(f'{where}: {code!r} -> {new!r} does not read back the same')
    return new, True


def packed(obj):                                                    # as tools/learned_to_js.py writes the parts
    return base64.b64encode(gzip.compress(json.dumps(obj, separators=(',', ':'), ensure_ascii=False).encode('utf-8'), 9)).decode('ascii')


def workbook(check):
    from openpyxl import load_workbook
    path = os.path.join(APP, 'reaction_rules.xlsx')
    wb = load_workbook(path)
    ws = wb['Reactions']
    col = [c.value for c in ws[1]].index('Reaction code') + 1
    n = 0
    for row in range(2, ws.max_row + 1):
        c = ws.cell(row, col)
        new, changed = convert(c.value, f'reaction_rules.xlsx Reactions!{c.coordinate}')
        if changed: n += 1; c.value = new
    how = wb['How to use']
    for row in how.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('Reaction code:') and c.value != HOW_TO: c.value = HOW_TO; n += 1
    if n and not check: wb.save(path)
    return n


def tests_json(name, check):
    path = os.path.join(HERE, name)
    if not os.path.exists(path): return 0
    T = json.load(open(path, encoding='utf-8'))
    n = 0
    for group, cases in T.items():
        if not isinstance(cases, list): continue
        for i, case in enumerate(cases):
            if not isinstance(case, dict) or 'code' not in case or case.get('name') in KEEP_V1: continue
            new, changed = convert(case['code'], f'{name} {group}[{i}]')
            if changed: n += 1; case['code'] = new
    if n and not check:
        with open(path, 'w', encoding='utf-8') as f: f.write(json.dumps(T, ensure_ascii=False, indent=1))
    return n


def learned(check):
    d = os.path.join(APP, 'data', 'learned', 'rx')
    if not os.path.isdir(d): return 0
    total = 0
    for fn in sorted(os.listdir(d)):
        path = os.path.join(d, fn)
        s = open(path, encoding='utf-8').read()
        m = re.match(r'(\(window\.LEARNED_PARTS=window\.LEARNED_PARTS\|\|\{\}\)\["[^"]+"\]=)\{"gz":"([^"]+)"\};\n?$', s)
        if not m: raise ValueError(f'{path}: not a learned part')
        rows = json.loads(gzip.decompress(base64.b64decode(m.group(2))))
        n = 0
        for r in rows:
            new, changed = convert(r[1], f'{fn} {r[0]}')
            if changed: n += 1; r[1] = new
        if n and not check:
            with open(path, 'w', encoding='utf-8') as f: f.write(f'{m.group(1)}{{"gz":"{packed(rows)}"}};\n')
        total += n
    return total


def main():
    check = '--check' in sys.argv
    verb = 'still version 1' if check else 'rewritten as version 2'
    for what, fn in (('reaction_rules.xlsx', lambda: workbook(check)),
                     ('tools/reaction_tests.json', lambda: tests_json('reaction_tests.json', check)),
                     ('tools/learned_tests.json', lambda: tests_json('learned_tests.json', check)),
                     ('data/learned/rx', lambda: learned(check))):
        print(f'{what}: {fn():,} {verb}', flush=True)
    if not check: print('now run: python tools/rules_to_js.py   (data/reaction_rules.js from the workbook)')


if __name__ == '__main__':
    main()
