"""Fragment IDs are generated from the SMARTS pattern, never chosen by hand:  F_ + 8 base32 characters of SHA-1(SMARTS).

The same pattern always gives the same ID, so two fragments with one pattern are caught at once.  The human label is
the 'Fragment name' column.  Editing a fragment's SMARTS changes its ID; after editing the Fragments sheet run

    python tools/fragment_id.py          # list IDs that no longer match their SMARTS
    python tools/fragment_id.py --fix    # re-ID those fragments everywhere (workbooks, JS, learned data, tools)
"""
import base64, glob, gzip, hashlib, json, os, re, sys, zipfile
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ID_RE = re.compile(r"F_[A-Z2-7]{8}")
TEXT_FILES = ['tools/data/fragments_learned.json', 'tools/pka.py', 'tools/learned_molecules.py', 'js/reactions.js', 'js/molview.js', 'data/reaction_rules.js']
BOOKS = ['molecule_data.xlsx', 'reaction_rules.xlsx']

def fragment_id(smarts):
    return "F_" + base64.b32encode(hashlib.sha1(str(smarts).strip().encode("utf-8")).digest()).decode()[:8]

def _sub(mapping, s):
    return re.sub(r"\bF_[A-Z0-9_]+\b", lambda m: mapping.get(m.group(0), m.group(0)), s)

def apply_map(mapping):
    """rename fragment IDs (old -> new) in every file that stores them"""
    p = lambda f: os.path.join(ROOT, f)
    for f in TEXT_FILES:
        s = open(p(f), encoding="utf8").read(); open(p(f), "w", encoding="utf8").write(_sub(mapping, s))
    for f in BOOKS:                                   # only inside <t> text nodes of inline strings and sharedStrings.xml (cell refs stay alone)
        tmp = p(f) + ".tmp"
        with zipfile.ZipFile(p(f)) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
            for it in zi.infolist():
                d = zi.read(it.filename)
                if it.filename.startswith("xl/worksheets/") or it.filename == "xl/sharedStrings.xml":
                    d = re.sub(r"(<t(?: [^>]*)?>)([^<]*)(</t>)", lambda m: m.group(1) + _sub(mapping, m.group(2)) + m.group(3), d.decode("utf8")).encode("utf8")
                zo.writestr(it, d)
        os.replace(tmp, p(f))
    for f in glob.glob(p("data/learned/rx/*.js")):    # columns 3, 4 = fragments consumed / formed
        s = open(f, encoding="utf8").read(); mm = re.search(r'"gz":"([^"]+)"', s)
        rows = json.loads(gzip.decompress(base64.b64decode(mm.group(1))).decode("utf8"))
        for r in rows:
            for i in (3, 4): r[i] = _sub(mapping, r[i])
        b = base64.b64encode(gzip.compress(json.dumps(rows, separators=(",", ":"), ensure_ascii=False).encode("utf8"), 9)).decode("ascii")
        open(f, "w", encoding="utf8").write(s[:mm.start(1)] + b + s[mm.end(1):])

DATA_BY_ID = ['tools/data/fragment_models.json', 'tools/data/fragment_properties.json', 'tools/data/fragment_conditions.json']


def forget(ids):
    """drop the generated rows (model compound, properties, conditions) of fragments whose SMARTS, and so ID, changed: they described the
    old pattern. tools/fragment_properties.py and tools/fragment_conditions.py then compute the new ones."""
    for f in DATA_BY_ID:
        p = os.path.join(ROOT, f)
        if not os.path.exists(p): continue
        d = json.load(open(p, encoding="utf-8"))
        if isinstance(d, list): d = [r for r in d if r.get("Fragment ID") not in ids]
        else: d = {k: v for k, v in d.items() if k not in ids}
        json.dump(d, open(p, "w", encoding="utf-8"), indent=1, ensure_ascii=False)


def sheet_rows():
    from openpyxl import load_workbook
    it = load_workbook(os.path.join(ROOT, "molecule_data.xlsx"), read_only=True)["Fragments"].iter_rows(values_only=True)
    head = [str(h) for h in next(it)]
    return [dict(zip(head, r)) for r in it if r and r[0]]

def stale():
    return {str(f["Fragment ID"]).strip(): fragment_id(f["SMARTS pattern"]) for f in sheet_rows()
            if str(f["Fragment ID"]).strip() != fragment_id(f["SMARTS pattern"])}

if __name__ == "__main__":
    m = stale()
    for a, b in m.items(): print(a, "->", b)
    if "--fix" in sys.argv and m:
        apply_map(m); forget(set(m) | set(m.values()))
        path = os.path.join(ROOT, "tools", "data", "fragment_ids.json"); old = json.load(open(path, encoding="utf8"))
        json.dump({k: m.get(v, v) for k, v in old.items()}, open(path, "w", encoding="utf8"), indent=1, ensure_ascii=False)
        print("re-ID'd", len(m), "fragment(s); run tools/rules_to_js.py, tools/fragment_properties.py --resume (computes only the re-ID'd fragments), tools/fragment_conditions.py, tools/fragment_properties_sheet.py and tools/run_all_checks.py")
    elif m: print(len(m), "ID(s) out of date; run with --fix"); sys.exit(1)
    else: print("all fragment IDs match their SMARTS")
