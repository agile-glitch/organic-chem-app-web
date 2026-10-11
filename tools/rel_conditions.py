"""The recorded solvents of every learned reaction (the 'Conditions (recorded solvents)' column of learned_reactions.xlsx), in Learned-ID order.

    python tools/rel_conditions.py            # about 10 minutes; writes organic_chem_rel_conditions.pkl in the temp folder

tools/rel_features.py --solvents reads it to add one bit per solvent (the Solvents sheet of reaction_rules.xlsx: 39 names, the same names the
Reactions page gets from the conditions box) to the model's input. Kept apart from rel_common.read_rows so the 1.09 M cached rows stay as they are.
"""
import os, pickle, sys, tempfile, time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import rel_common as C

OUT = os.path.join(tempfile.gettempdir(), 'organic_chem_rel_conditions.pkl')


def read():
    import openpyxl
    t0 = time.time()
    wb = openpyxl.load_workbook(C.workbook_path(), read_only=True, data_only=True)
    ids, conds = [], []
    for name in wb.sheetnames:
        if not name.startswith('Learned'): continue
        head = None
        for r in wb[name].iter_rows(values_only=True):
            if head is None: head = {h: k for k, h in enumerate(r)}; continue
            if not r or not r[0]: continue
            ids.append(r[head['Learned ID']]); conds.append(r[head['Conditions (recorded solvents)']] or '')
        print(f'  {name}: {len(ids):,} ({time.time() - t0:.0f} s)', flush=True)
    pickle.dump({'ids': ids, 'conds': conds}, open(OUT, 'wb'), protocol=4)
    print(f'wrote {OUT}: {len(ids):,} reactions, {sum(1 for c in conds if c) / len(conds):.1%} with a recorded solvent')


def load():
    """[solvent string per reaction], checked against the cached rows"""
    d = pickle.load(open(OUT, 'rb'))
    rows = C.read_rows()['rows']
    assert len(rows) == len(d['ids']) and rows[0][0] == d['ids'][0] and rows[-1][0] == d['ids'][-1], 'conditions do not match the cached rows'
    return d['conds']


if __name__ == '__main__':
    read()
