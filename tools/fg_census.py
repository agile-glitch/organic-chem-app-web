"""Functional-group census: in how many fragments / interference rows / templates does each group occur (entries and occurrences)?
Also the average count of each fragment per recorded reaction. Writes data/fragment_census.js (the Validation tab reads it) and the json the Word report is built from.

    python tools/fg_census.py        # then: python tools/make_report_docx.py
"""
import sys, os, re, json, tempfile, collections
import numpy as np
REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, 'tools'))
from rdkit import Chem, RDLogger
RDLogger.DisableLog('rdApp.*')
import interference_defs as D
from openpyxl import load_workbook

# ---------------------------------------------------------------- functional groups (SMARTS written for a concrete molecule)
FG = [
 ('Alcohol (sp3 C–OH)', '[CX4][OX2H]'), ('Phenol (Ar–OH)', 'c[OX2H]'), ('Carboxylic acid', '[CX3](=O)[OX2H1]'), ('Carboxylate anion', '[CX3](=O)[OX1-]'),
 ('Ester / lactone', '[#6][CX3](=O)[OX2][#6]'), ('Amide / lactam', '[#6][CX3](=O)[NX3]'), ('Carbamate', '[NX3][CX3](=O)[OX2]'), ('Urea', '[NX3][CX3](=O)[NX3]'),
 ('Acyl halide', '[CX3](=O)[F,Cl,Br,I]'), ('Anhydride', '[CX3](=O)[OX2][CX3](=O)'), ('Aldehyde', '[CX3H1](=O)[#6]'), ('Ketone', '[#6][CX3](=O)[#6]'),
 ('Ether (C–O–C, not ester)', '[#6;!$(C=O)][OX2][#6;!$(C=O)]'), ('Epoxide', 'C1OC1'), ('Acetal / ketal', '[CX4]([OX2])[OX2]'), ('Enol ether', 'C=C[OX2][#6]'),
 ('Primary amine', '[NX3;H2][#6]'), ('Secondary amine', '[NX3;H1]([#6])[#6]'), ('Tertiary amine', '[NX3;H0]([#6])([#6])[#6]'), ('Aniline-type N (N on aromatic C)', '[NX3]c'),
 ('Ammonium', '[NX4+]'), ('Imine', '[CX3]=[NX2]'), ('Nitrile', 'C#N'), ('Nitro', '[N+](=O)[O-]'), ('Azide', 'N=[N+]=[N-]'), ('Hydrazine / hydrazide', '[NX3][NX3]'),
 ('Oxime / hydroxylamine', '[#7][OX2H]'), ('Diazonium / diazo', '[N+]#N'), ('Isocyanate / isothiocyanate', 'N=C=[O,S]'),
 ('Thiol', '[SX2H]'), ('Thioether', '[#6][SX2][#6]'), ('Sulfoxide', '[SX3](=O)'), ('Sulfone', '[SX4](=O)(=O)([#6])[#6]'), ('Sulfonamide', '[SX4](=O)(=O)[NX3]'),
 ('Sulfonate / sulfonyl halide', '[SX4](=O)(=O)[OX2,F,Cl]'),
 ('Alkyl fluoride', '[CX4]F'), ('Alkyl chloride', '[CX4]Cl'), ('Alkyl bromide', '[CX4]Br'), ('Alkyl iodide', '[CX4]I'), ('Trifluoromethyl', 'C(F)(F)F'),
 ('Aryl fluoride', 'cF'), ('Aryl chloride', 'cCl'), ('Aryl bromide', 'cBr'), ('Aryl iodide', 'cI'), ('Vinyl halide', 'C=C[Cl,Br,I]'),
 ('Alkene (C=C)', '[CX3]=[CX3]'), ('Alkyne (C≡C)', 'C#C'), ('Conjugated carbonyl (enone / acrylate)', 'C=C[CX3]=O'), ('Benzene ring (aromatic carbon)', 'c'),
 ('Pyridine-type ring N', '[nX2]'), ('Pyrrole-type ring N', '[nX3]'), ('Furan / thiophene / oxazole ring O,S', '[o,s]'),
 ('Phosphorus (phosphonate, phosphine, phosphate)', '[P]'), ('Boron (boronic acid / ester)', '[B]'), ('Silicon (silyl)', '[Si]'),
 ('Lithium / magnesium organometallic', '[Li,Mg]'), ('Palladium / nickel / copper / other metal', '[Pd,Ni,Cu,Zn,Sn,Fe,Pt,Rh,Ru,Ir,Co,Mn,Cr,Ti,Zr,Al,Ag,Hg]'),
 ('Carbamate protecting group (Boc / Cbz-type)', '[NX3][CX3](=O)[OX2][CX4,CH2]'), ('tert-Butyl group', '[CX4]([CH3])([CH3])[CH3]'), ('Methyl group (on C)', '[CH3][#6]'),
]
FGQ = [(n, Chem.MolFromSmarts(s)) for n, s in FG]
assert all(q is not None for _, q in FGQ)

_H = re.compile(r'H(\d?)')

def skeleton(sma):
    """a concrete molecule drawn from a SMARTS: element, charge, aromaticity, bond order and an explicit H count when the SMARTS gives one"""
    q = Chem.MolFromSmarts(sma)
    if q is None: return None
    rw = Chem.RWMol()
    for a in q.GetAtoms():
        z = a.GetAtomicNum() or 6
        at = Chem.Atom(z)
        sm = a.GetSmarts()
        ch = a.GetFormalCharge()
        if ch == 0:
            mch = re.search(r'([+-])(\d?)', sm.split(':')[0])
            if mch and not re.search(r'[+-]0', sm): ch = (1 if mch.group(1) == '+' else -1) * (int(mch.group(2)) if mch.group(2) else 1)
        at.SetFormalCharge(ch)
        at.SetIsAromatic(a.GetIsAromatic())
        mh = re.search(r';?H(\d?)(?![a-z])', sm) if z != 1 else None
        if mh and '$' not in sm.split('H')[0][-3:]:
            at.SetNumExplicitHs(int(mh.group(1) or 1)); at.SetNoImplicit(True)
        rw.AddAtom(at)
    for b in q.GetBonds():
        i, j = b.GetBeginAtomIdx(), b.GetEndAtomIdx()
        bt = b.GetBondType()
        arom = b.GetIsAromatic() or (bt == Chem.BondType.AROMATIC)
        if not arom and q.GetAtomWithIdx(i).GetIsAromatic() and q.GetAtomWithIdx(j).GetIsAromatic() and b.GetSmarts() in ('', '~'): arom = True
        if arom: t = Chem.BondType.AROMATIC
        elif bt in (Chem.BondType.DOUBLE, Chem.BondType.TRIPLE): t = bt
        else: t = Chem.BondType.SINGLE
        rw.AddBond(i, j, t)
        if arom: rw.GetBondBetweenAtoms(i, j).SetIsAromatic(True)
    m = rw.GetMol()
    m.UpdatePropertyCache(strict=False)
    Chem.FastFindRings(m)
    return m


def groups_in(sma):
    m = skeleton(sma)
    if m is None: return collections.Counter()
    out = collections.Counter()
    for n, q in FGQ:
        try:
            k = len(m.GetSubstructMatches(q, uniquify=True, maxMatches=500))
            if k: out[n] = k
        except Exception: pass
    return out


def strip_negations_and_split(sma):
    """a site pattern with recursion: the main pattern with the $() removed, plus each positive $() body (negated ones are what the atom must NOT be)"""
    parts, main, i = [], [], 0
    while i < len(sma):
        if sma.startswith('$(', i):
            neg = i > 0 and sma[i - 1] == '!'
            depth, j = 1, i + 2
            while j < len(sma) and depth:
                depth += (sma[j] == '(') - (sma[j] == ')'); j += 1
            body = sma[i + 2:j - 1]
            if not neg: parts.append(body)
            i = j
        else:
            main.append(sma[i]); i += 1
    cleaned = ''.join(main).replace('!', '')
    cleaned = re.sub(r';+\]', ']', cleaned); cleaned = re.sub(r'\[;', '[', cleaned); cleaned = re.sub(r';;+', ';', cleaned)
    return [cleaned] + parts

# ---------------------------------------------------------------- fragments
wb = load_workbook(os.path.join(REPO, 'molecule_data.xlsx'), read_only=True)
frag = {r[0]: (r[1], r[2]) for r in wb['Fragments'].iter_rows(min_row=2, values_only=True) if r[0]}
frag_groups = {fid: groups_in(sma) for fid, (_, sma) in frag.items()}
# ---------------------------------------------------------------- interference rows
site_groups = {}
for s in D.SITE:
    g = collections.Counter()
    for part in strip_negations_and_split(s['smarts']):
        g.update(groups_in(part))
    site_groups[s['key']] = g
adv_groups = {}
for a in D.ADVISORY:
    g = collections.Counter()
    for f in a['groups']: g.update(frag_groups.get(f, {}))
    adv_groups[a['key']] = g
# ---------------------------------------------------------------- templates (the 39,110 classes of the installed network)
import int_export as X
Z = np.load(os.path.join(tempfile.gettempdir(), 'organic_chem_relmodel_full2.npz'))
classes = [int(c) for c in Z['classes']]
tsm = X.templates_smarts()
cnt_r, cnt_p, cnt_any = collections.Counter(), collections.Counter(), collections.Counter()
occ_r, occ_p = collections.Counter(), collections.Counter()
ncls = 0
cache = {}
for t in classes:
    s = tsm[t]
    if '>>' not in s: continue
    left, right = s.split('>>', 1)              # forward templates: reactants>>products
    for side, c in ((left, cnt_r), (right, cnt_p)):
        if side not in cache: cache[side] = groups_in(side)
    gr, gp = cache[left], cache[right]
    ncls += 1
    for g, k in gr.items(): cnt_r[g] += 1; occ_r[g] += k
    for g, k in gp.items(): cnt_p[g] += 1; occ_p[g] += k
    for g in set(gr) | set(gp): cnt_any[g] += 1
res = {'n_fragments': len(frag), 'n_site': len(D.SITE), 'n_adv': len(D.ADVISORY), 'n_templates': ncls,
       'rows': []}
for n, _ in FG:
    f = sum(1 for g in frag_groups.values() if n in g); fo = sum(g[n] for g in frag_groups.values())
    si = sum(1 for g in site_groups.values() if n in g); sio = sum(g[n] for g in site_groups.values())
    ad = sum(1 for g in adv_groups.values() if n in g); ado = sum(g[n] for g in adv_groups.values())
    res['rows'].append({'group': n, 'fragments': f, 'site': si, 'adv': ad, 'interferences': si + ad, 'both': f + si + ad, 'tpl_reactant': cnt_r[n], 'tpl_product': cnt_p[n], 'tpl_any': cnt_any[n],
                        'o_fragments': fo, 'o_interferences': sio + ado, 'o_both': fo + sio + ado, 'o_tpl_reactant': occ_r[n], 'o_tpl_product': occ_p[n], 'o_tpl_any': occ_r[n] + occ_p[n]})
json.dump(res, open(os.path.join(tempfile.gettempdir(), 'fg_census.json'), 'w'), indent=1)
zc = np.load(os.path.join(tempfile.gettempdir(), 'organic_chem_relfragcounts_new2.npz'), allow_pickle=True)
SC = zc['S']; frs = [str(x) for x in zc['fragments']]
ma = SC.mean(0); sh = (SC >= 1).mean(0)
res['reactions'] = int(SC.shape[0]); res['hits_per_reaction'] = round(float(SC.sum(1).mean()), 2); res['distinct_per_reaction'] = round(float((SC >= 1).sum(1).mean()), 2)
res['fragment_stats'] = [[f, frag.get(f, ('', ''))[0], round(float(ma[i]), 4), round(float(sh[i]), 5), round(float(ma[i] / sh[i]), 3) if sh[i] > 0 else 0] for i, f in enumerate(frs)]
with open(os.path.join(REPO, 'data', 'fragment_census.js'), 'w', encoding='utf-8') as fh:
    fh.write('/* generated by tools/fg_census.py: functional groups in the fragments, interference rows and templates; fragment averages per recorded reaction */\nwindow.FRAGMENT_CENSUS = ' + json.dumps(res, ensure_ascii=False) + ';\n')
print('wrote data/fragment_census.js')
print(ncls, 'templates;', len(frag), 'fragments;', len(D.SITE), '+', len(D.ADVISORY), 'interference rows')
for r in sorted(res['rows'], key=lambda r: -r['tpl_any'])[:60]:
    print(f"{r['group'][:44]:44s} frag {r['fragments']:3d} int {r['interferences']:3d} both {r['both']:3d} | tpl react {r['tpl_reactant']:6d} prod {r['tpl_product']:6d} any {r['tpl_any']:6d}")
