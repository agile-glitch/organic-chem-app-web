"""Build learned_molecules.xlsx: one row per distinct molecule of the learned reactions, with its fragments and how
each fragment behaves in that molecule in the recorded reactions.

    python tools/learned_molecules.py

Reads learned_reactions.xlsx (tools/learn_reactions.py), the Fragments sheet of molecule_data.xlsx and the reagent
names of tools/data/reagents.json. For every distinct starting material and product:
  * SMILES, formula, molecular weight, monoisotopic mass, InChI, InChIKey (RDKit), and whether the page can read it;
  * how many learned reactions use it as a starting material / make it as the product, with examples;
  * its fragments (each with how many times it occurs);
  * how each fragment behaves in THIS molecule, from the recorded reactions:
      as a starting material - reacted in n of m reactions (an atom of the fragment made, broke or changed a bond:
      the "Fragments that reacted" column of learned_reactions.xlsx), with the commonest reagents; left untouched in
      the others, with the commonest reagents it survived;
      as a product - formed in n reactions (the fragment is new: "Fragments formed"), carried through unchanged in the rest.
Molecules are ordered by how many learned reactions use them (M0000001 = the most used).
"""
import collections, json, os, re, sys, time
from multiprocessing import Pool

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
sys.path.insert(0, HERE)
SRC = os.path.join(APP, 'learned_reactions.xlsx')
OUT = os.path.join(APP, 'learned_molecules.xlsx')
ROWS_PER_SHEET = 250000
REACTED_COL = 'Fragments that reacted, per starting material (in code order, | between molecules)'

_P = {}


def _init(frags, pka_defaults, pka_measured):
    from rdkit import Chem, RDLogger
    RDLogger.DisableLog('rdApp.*')
    _P['frags'] = [(fid, Chem.MolFromSmarts(sm)) for fid, _, sm in frags]
    _P['pka'] = (pka_defaults, pka_measured)


def props(smiles):
    """formula, MW, monoisotopic mass, InChI, InChIKey, page readable, {fragment: occurrences}"""
    import learn_reactions as L
    from rdkit import Chem
    from rdkit.Chem import Descriptors, rdMolDescriptors
    m = Chem.MolFromSmiles(smiles)
    if m is None: return smiles, None
    try: inchi = Chem.MolToInchi(m) or ''
    except Exception: inchi = ''
    key = Chem.InchiToInchiKey(inchi) if inchi else ''
    fr = {}
    for fid, q in _P['frags']:
        if q is not None:
            n = len(m.GetSubstructMatches(q, uniquify=True, maxMatches=1000))
            if n: fr[fid] = n
    import pka as PK
    try:
        sites = PK.molecule_pkas(smiles, _P['frags'], _P['pka'][0], _P['pka'][1].get(smiles))
        pk = '; '.join(f'{"/".join(sorted(set(x["fragments"])))} {"pKa" if kind == "acid" else "pKaH"} {x["pKa"]} ({x["source"]})'
                       for kind in ('acid', 'base') for x in sites[kind])
    except Exception as e:
        pk = ''
    import effects as EF
    try:                                                 # steric and electronic effects at each fragment's reacting atom
        E = [e for e in EF.effects_of(m, _P['frags']) if e['fragment'] != 'F009']
        st = '; '.join(f'{e["fragment"]} {e["steric"]["text"]}' for e in E)
        el = '; '.join(f'{e["fragment"]} {e["electronic"]["text"]}' for e in E)
        tf = '; '.join(f'{e["fragment"]} ' + ', '.join(x for x in (
            f'Σσ* {e["taft"]["sigma_star"]:+.2f}' if e['taft']['sigma_star'] is not None else '',
            f'bulkiest Es {e["taft"]["Es_min"]:+.2f}' if e['taft']['Es_min'] is not None else '',
            ' / '.join(g['name'] for g in e['taft']['groups'])) if x) for e in E if e['taft']['groups'])
    except Exception:
        st = el = tf = ''
    return smiles, (rdMolDescriptors.CalcMolFormula(m), round(Descriptors.MolWt(m), 3), round(Descriptors.ExactMolWt(m), 4),
                    inchi, key or '', L.page_readable(smiles), fr, pk, st, el, tf)


def main():
    import argparse, openpyxl, learn_reactions as L
    global SRC, OUT
    ap = argparse.ArgumentParser(); ap.add_argument('--src', default=SRC); ap.add_argument('--out', default=OUT)
    a = ap.parse_args(); SRC, OUT = a.src, a.out
    t0 = time.time()
    frags = L.fragments()
    fname = {fid: name for fid, name, _ in frags}
    roles = {x['smiles']: x for x in json.load(open(os.path.join(HERE, 'data', 'reagents.json'), encoding='utf-8'))}
    label_of, labels = {}, []

    def label(reagent):                                  # a short reagent name: "HATU + DIPEA" (spectator ions left out)
        if reagent not in label_of:
            names = []
            for p in L.reagent_counted(reagent):
                r = roles.get(p)
                if r and r.get('role') == 'counter-ion / spectator ion': continue
                n = (r.get('abbreviation') or r.get('name')) if r else p
                n = re.split(r'[ (]', n.strip())[0] if r and r.get('abbreviation') else n
                names.append(n)
            label_of[reagent] = len(labels); labels.append(' + '.join(dict.fromkeys(names)) or ('no reagent' if not reagent else reagent))
        return label_of[reagent]

    # ---- read the learned reactions ----
    mol_id, mols = {}, []
    def mid(s):
        if s not in mol_id: mol_id[s] = len(mols); mols.append(s)
        return mol_id[s]
    as_sm, as_pr = collections.Counter(), collections.Counter()
    ex_sm, ex_pr = collections.defaultdict(list), collections.defaultdict(list)
    sm_labels = collections.defaultdict(list)             # molecule -> reagent label of every reaction it starts
    reacted = collections.defaultdict(list)               # (molecule, fragment) -> reagent label of every reaction it reacted in
    formed = collections.Counter()                        # (product, fragment) -> times formed
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    n = 0
    for name in wb.sheetnames:
        if not name.startswith('Learned'): continue
        head = None
        for r in wb[name].iter_rows(values_only=True):
            if head is None: head = {h: k for k, h in enumerate(r)}; continue
            if not r or not r[0]: continue
            n += 1
            lid = r[head['Learned ID']]
            subs = (r[head['Starting materials']] or '').split('.')
            reagent, prod = r[head['Reagent']] or '', r[head['Product']]
            lab = label(reagent)
            rx = (r[head[REACTED_COL]] or '').split(' | ') if REACTED_COL in head else []
            for j, s in enumerate(subs):
                if not s: continue
                k = mid(s); as_sm[k] += 1; sm_labels[k].append(lab)
                if len(ex_sm[k]) < 3: ex_sm[k].append(lid)
                for fid in (rx[j] if j < len(rx) else '').split(' + '):
                    if fid and fid != '-': reacted[(k, fid)].append(lab)
            k = mid(prod); as_pr[k] += 1
            if len(ex_pr[k]) < 3: ex_pr[k].append(lid)
            for fid in (r[head['Fragments formed']] or '').split(' + '):
                if fid: formed[(k, fid)] += 1
            if n % 200000 == 0: print(f'  read {n:,} reactions ({time.time() - t0:.0f} s)', flush=True)
    print(f'{n:,} learned reactions, {len(mols):,} distinct molecules ({time.time() - t0:.0f} s)', flush=True)

    # ---- molecule properties and fragments (in parallel) ----
    P = {}
    import pka as PK
    pka_measured = json.load(open(os.path.join(HERE, 'data', 'pka_measured.json'), encoding='utf-8'))['molecules']
    with Pool(max(1, (os.cpu_count() or 2) - 1), initializer=_init, initargs=(frags, PK.fragment_defaults(), pka_measured)) as pool:
        for i, (s, pr) in enumerate(pool.imap_unordered(props, mols, chunksize=500)):
            P[mol_id[s]] = pr
            if (i + 1) % 200000 == 0: print(f'  described {i + 1:,} molecules ({time.time() - t0:.0f} s)', flush=True)
    print(f'described {len(P):,} molecules ({time.time() - t0:.0f} s)', flush=True)

    # ---- rows ----
    def top(lst, k=3):
        return ', '.join(f'{labels[l]} ×{c}' for l, c in collections.Counter(lst).most_common(k))
    order = sorted(range(len(mols)), key=lambda k: (-(as_sm[k] + as_pr[k]), mols[k]))
    rows = []
    for rank, k in enumerate(order, 1):
        pr = P.get(k)
        if pr is None: continue
        formula, mw, mono, inchi, key, readable, fr, pk, st, el, tf = pr
        frag_text = '; '.join(f'{fid} {fname.get(fid, "")} ×{c}' for fid, c in sorted(fr.items()))
        beh = []
        for fid in sorted(fr):
            parts = []
            if as_sm[k]:
                rl = reacted.get((k, fid), [])
                if rl:
                    parts.append(f'as a starting material reacted in {len(rl)} of {as_sm[k]} ({top(rl)})')
                    left = collections.Counter(sm_labels[k]); left.subtract(collections.Counter(rl))
                    left = +left
                    if sum(left.values()): parts.append(f'left untouched in {sum(left.values())} ({", ".join(f"{labels[l]} ×{c}" for l, c in left.most_common(3))})')
                else:
                    parts.append(f'as a starting material never reacted: left untouched in all {as_sm[k]} ({top(sm_labels[k])})')
            if as_pr[k]:
                f = formed.get((k, fid), 0)
                parts.append(f'as a product formed in {f} of {as_pr[k]}' + (f', carried through unchanged in {as_pr[k] - f}' if as_pr[k] - f else '') if f
                             else f'as a product carried through unchanged in all {as_pr[k]}')
            beh.append(f'{fid} {fname.get(fid, "")}: ' + '; '.join(parts))
        rows.append([f'M{rank:07d}', mols[k], formula, mw, mono, inchi, key, 'yes' if readable else 'no', as_sm[k], as_pr[k],
                     frag_text, pk, st, el, tf, ' | '.join(beh), ', '.join(ex_sm[k]), ', '.join(ex_pr[k])])
    print(f'{len(rows):,} molecule rows ({time.time() - t0:.0f} s)', flush=True)
    write(rows, frags, n)
    print(f'wrote {OUT} ({time.time() - t0:.0f} s)')


def write(rows, frags, n_rx):
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    HDR, FILL = Font(name='Arial', bold=True, color='FFFFFF'), PatternFill('solid', fgColor='1F4E78')
    wb = Workbook(write_only=True)
    col = lambda k: (chr(64 + k // 26) if k >= 26 else '') + chr(65 + k % 26)

    def sheet(title, head, widths, data):
        ws = wb.create_sheet(title)
        for k, w in enumerate(widths): ws.column_dimensions[col(k)].width = w
        ws.freeze_panes = 'C2'
        cells = []
        for h in head:
            c = WriteOnlyCell(ws, value=h); c.font, c.fill = HDR, FILL; c.alignment = Alignment(wrap_text=True, vertical='center'); cells.append(c)
        ws.append(cells)
        for r in data: ws.append(r)

    nsheets = (len(rows) - 1) // ROWS_PER_SHEET + 1
    sheet('How to use', ['How this workbook is organised'], [140], [[t] for t in [
        f'Molecules 1-{nsheets}: the {len(rows):,} distinct molecules of the {n_rx:,} learned reactions (learned_reactions.xlsx), one row each, '
        f'{ROWS_PER_SHEET:,} rows per sheet, the most used first (M0000001).',
        'SMILES (RDKit canonical, the spelling the reaction codes use), formula, molecular weight (average, g/mol), monoisotopic mass, InChI and '
        'InChIKey: all computed by RDKit from the SMILES. Names are not given (they would need a look-up per molecule).',
        'Fragments: every fragment of the Fragments sheet (molecule_data.xlsx) found in the molecule, with how many times it occurs.',
        'How each fragment behaves: from the recorded reactions. As a starting material, a fragment "reacted" when one of its atoms made, broke or '
        'changed a bond (every atom is followed from the starting materials into the product by the reaction template); otherwise it was "left '
        'untouched". The commonest reagents are listed for each (counter-ions left out). As a product, a fragment was "formed" when the reaction '
        'made it new, else it was carried through unchanged.',
        'pKa in this molecule: each acidic and basic site (one atom) gets the measured value of the molecule if it has one (IUPAC Digitized pKa '
        'Dataset v2.4, CC BY-NC 4.0, reproduced by permission of IUPAC; water, 20-30 °C), else a Hammett estimate for substituted phenols, '
        'benzoic acids and anilines (sigma values: Hansch, Leo, Taft, Chem. Rev. 1991), else the typical value of the fragment, and says which.',
        'Steric and electronic effects: read at the reacting atom of each fragment (the C of a C=O, a C carrying a halogen, an N/O/S, a C=C carbon). '
        'Steric score = its heavy neighbours + the most further heavy neighbours on one of them (aromatic: ortho substituents only): <= 2 open, '
        '3 moderately hindered, 4 hindered, >= 5 very hindered. Electronic: Gasteiger partial charge, the meta/para Hammett sigma of a benzene '
        'ring it sits on or next to (sum > 0 electron-poor, < 0 electron-rich), conjugation, and the oxidation state of a carbon (tools/effects.py).',
        'Taft parameters: each group R on the reacting atom (not its own ring, not a =O/=S partner) with its sigma* (polar; CH3 = 0, > 0 pulls '
        'electrons, < 0 pushes) and Es (steric; CH3 = 0, more negative = bulkier). Values: Taft 1956; Hansch, Leo, Taft, Chem. Rev. 1991; '
        'Hansch & Leo 1979. A carbon group not in the table gets an Es estimated from its branching and no sigma*.',
        'What a fragment does in general (role, typical pKa, leaving group, steric hindrance): the Fragments sheet of this workbook.',
        'Rebuild: python tools/learned_molecules.py (after python tools/learn_reactions.py).']])
    import learn_reactions as L
    fsheet = L.read_sheet(os.path.join(APP, 'molecule_data.xlsx'), 'Fragments')
    keep = ['Fragment ID', 'Fragment name', 'SMARTS pattern', 'Reactive atom(s)', 'Role', 'pKa (its most acidic H)',
            'pKa of its conjugate acid (how basic)', 'Leaving group', 'Steric hindrance (at the reactive atom)', 'Notes']
    sheet('Fragments', keep, [11, 30, 40, 24, 40, 22, 22, 30, 34, 50], [[f.get(k) for k in keep] for f in fsheet if str(f.get('Fragment ID', '')).startswith('F')])
    head = ['Molecule ID', 'SMILES', 'Formula', 'Molecular weight (g/mol)', 'Monoisotopic mass', 'InChI', 'InChIKey', 'Page can read it',
            'Starting material in (learned reactions)', 'Product of (learned reactions)', 'Fragments (ID name ×occurrences)',
            'pKa of its acidic (pKa) and basic (pKaH) sites in THIS molecule (measured, estimated or typical)',
            'Steric effects at the reacting atom of each fragment (open / moderately hindered / hindered / very hindered)',
            'Electronic effects at the reacting atom of each fragment (partial charge, ring sigma, conjugation, oxidation state)',
            'Taft parameters of the groups on the reacting atom of each fragment (sum of sigma*, bulkiest Es, the groups)',
            'How each fragment behaves in this molecule (recorded reactions)', 'Examples as a starting material', 'Examples as a product']
    widths = [11, 50, 16, 12, 12, 50, 30, 9, 12, 12, 50, 60, 60, 60, 50, 90, 26, 26]
    for k in range(0, len(rows), ROWS_PER_SHEET):
        sheet(f'Molecules {k // ROWS_PER_SHEET + 1}', head, widths, rows[k:k + ROWS_PER_SHEET])
    from learn_reactions import save_workbook
    save_workbook(wb, OUT)


if __name__ == '__main__':
    main()
