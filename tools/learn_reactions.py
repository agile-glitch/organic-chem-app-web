"""Learn reactions from the recorded-reaction knowledge base (data/rxnkb) into learned_reactions.xlsx.

    python tools/learn_reactions.py              # every recorded reaction that passes the check
    python tools/learn_reactions.py --n 20000    # only the first 20,000 (round-robin order, see below)

Where the reactions come from: data/rxnkb holds 1.1 million reactions recorded in US patents (1976-2016, text-mined
by D. M. Lowe, CC0) and the Open Reaction Database (CC-BY-SA 4.0), each with the reaction template rdchiral extracted
from it (the atoms that change + their neighbours, written as reaction SMARTS: the same kind of rule as the SMARTS
column of reaction_rules.xlsx, made by software instead of by hand).

The test every learned reaction passes (CHECKED): the template, run by RDKit on the starting materials exactly as
the page runs a learned template (every order, every site; the carbon-free reagent pieces such as OH-, Br2 or HNO3
may take part), must give the recorded product, stereo included. A reaction that fails is not learned and the reason
is counted. (This proves the row is consistent: template, molecules and product agree. It cannot prove the patent's
chemistry was right; text-mined records carry a few % of mistakes, and the check removes the ones that do not fit.)

Order (the Learned IDs): round-robin over the templates, commonest first, so L000001... cover as many DIFFERENT
reactions as possible before repeats; within a template, a different reagent first. The same reaction recorded
several times (same code) is learned once and counted.

What each learned reaction gets (one row of a "Learned n" sheet), in the same form as reaction_rules.xlsx:
  * the four-part reaction code: starting materials~reagent~conditions~product, molecules in the SMILES-line order;
    carbon-free reactants (Br2, HNO3, OH-, H2O ...) and the non-solvent agents are the reagent; the recorded solvents
    are the conditions (names from the Solvents sheet of reaction_rules.xlsx); recorded reactions have no temperature;
  * fragments consumed / formed, read with the Fragments sheet of molecule_data.xlsx (a fragment consumed = its count
    drops from the starting materials to the product; formed = its count rises);
  * leaving pieces: the starting-material atoms that are not in the product, with H added where bonds broke (the raw
    material for by-products, see the Byproducts sheet of reaction_rules.xlsx);
  * the template, how often it is recorded, the source (patent number), year, yield, how many times the same reaction
    was recorded, whether the page's own SMILES reader can read every molecule, and the look-up keys (InChIKeys).
Mechanisms: none. Recorded reactions give only starting materials -> product, never the arrows.

Then: python tools/learned_to_js.py (the page's data) and python tools/check_learned.py (re-checks every row).
"""
import argparse, base64, collections, gzip, itertools, json, os, re, sys, time
from multiprocessing import Pool

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
KB = os.path.join(APP, 'data', 'rxnkb')
OUT = os.path.join(APP, 'learned_reactions.xlsx')
ROWS_PER_SHEET = 250000
SAVED = os.path.join(__import__('tempfile').gettempdir(), 'organic_chem_learned_last_run.pkl')


def save_workbook(wb, path):
    """save; if the file is open in Excel (Windows locks it), save next to it as '<name> (new).xlsx' and say so"""
    try:
        wb.save(path); print('wrote', path); return path
    except PermissionError:
        alt = path[:-5] + ' (new).xlsx'
        wb.save(alt)
        print(f'*** {os.path.basename(path)} is open in another program (Excel?), so it could not be replaced. '
              f'Saved as {os.path.basename(alt)} instead: close the old one and rename this one, or run again with --write-only.')
        return alt

# the page's SMILES reader (js/chem.js): valences, metals without implicit H, "ate" formers, expanded octets
VALENCE = dict(C=4, N=3, O=2, S=2, P=3, F=1, Cl=1, Br=1, I=1, H=1, B=3, Si=4, Se=2, Sn=4, Li=1, Na=1, K=1, Mg=2, Ca=2,
               Zn=2, Cu=2, Fe=3, Al=3, Mn=7, Cr=6, Os=8, Pd=2, Pt=2, Hg=2, Pb=4, Ag=1, Cs=1, Rb=1, Ba=2, Sr=2, Ni=2,
               Co=2, Rh=3, Ru=3, Ir=3, Au=1, Ti=4, Zr=4, Ce=3, Yb=3, Sc=3, In=3, La=3, Sm=3, Bi=3, Mo=6, W=6, V=5, Cd=2,
               As=3, Sb=3, Ge=4, Te=2)
ATE = set('B Al Li Na K Mg Ca Zn Cu Fe Mn Cr Os Pd Pt Hg Pb Sn Ag Cs Rb Ba Sr Ni Co Rh Ru Ir Au Ti Zr Ce Yb Sc In La Sm Bi Mo W V Cd'.split())
EXPANDED = dict(S=6, P=6, Se=6, Si=4, Mn=7, Cr=6, Os=8, I=7, Cl=7, Br=5, As=6, Sb=6, Te=6, Pd=6, Pt=6, Ni=6, Co=6, Fe=6, Rh=6, Ru=6, Ir=6, Cu=4, Au=4, Ti=6,
                Zr=6, Ce=6, Yb=6, Sc=6, In=4, Mo=6, W=6, Sn=6)


# ---------- the knowledge base ----------
def _payload(path):
    s = open(path, encoding='utf-8').read()
    i = s.index('push(') + 5 if 'push(' in s else s.index('=', s.index(']')) + 1
    return json.loads(s[i:s.rindex('}') + 1])


def load_kb():
    import numpy as np
    meta = json.loads(open(os.path.join(KB, 'meta.js'), encoding='utf-8').read().split('=', 1)[1].rstrip(';\n '))
    tmpl, tn = [None] * meta['n_templates'], [0] * meta['n_templates']
    for f in meta['template_files']:
        p = _payload(os.path.join(KB, f))
        body = json.loads(gzip.decompress(base64.b64decode(p['gz']))) if 'gz' in p else p
        for j, (s, n) in enumerate(zip(body['smarts'], body['n'])):
            tmpl[p['start'] + j] = s; tn[p['start'] + j] = n
    tid = np.zeros(meta['n'], dtype=np.uint32)
    for f in meta['fp_files']:
        p = _payload(os.path.join(KB, f))
        raw = gzip.decompress(base64.b64decode(p['tid'])) if p.get('z') else base64.b64decode(p['tid'])
        a = np.frombuffer(raw, dtype=np.uint32); tid[p['start']:p['start'] + len(a)] = a
    R, P, A, C, Y, YL = [], [], [], [], [], []
    for k in range(meta['prec_files']):
        d = _payload(os.path.join(KB, f'prec_{k:04d}.js'))
        body = json.loads(gzip.decompress(base64.b64decode(d['gz'])))
        R += body['r']; P += body['p']; A += body['a']; C += body['c']; Y += body['y']; YL += body['yl']
    return dict(meta=meta, tmpl=tmpl, tn=tn, tid=tid.tolist(), r=R, p=P, a=A, c=C, y=Y, yl=YL)


# ---------- the workbooks ----------
def read_sheet(path, name):
    import openpyxl
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[name]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h or '').strip() for h in rows[0]]
    return [dict(zip(head, r)) for r in rows[1:] if r and r[0] not in (None, '')]


def solvents():
    from rdkit import Chem
    out = {}
    for s in read_sheet(os.path.join(APP, 'reaction_rules.xlsx'), 'Solvents'):
        m = Chem.MolFromSmiles(str(s.get('SMILES') or '')) if s.get('SMILES') else None
        if m: out[Chem.MolToSmiles(m)] = str(s['Solvent']).strip()
    return out


def fragments():
    return [(str(f['Fragment ID']).strip(), str(f['Fragment name']).strip(), str(f['SMARTS pattern']).strip())
            for f in read_sheet(os.path.join(APP, 'molecule_data.xlsx'), 'Fragments') if str(f['Fragment ID']).startswith('F')]


# ---------- the SMILES-line order (js/reactions.js sortByRules) ----------
def aromatic_atoms(smiles):                     # read from the text, as the page does
    n, i = 0, 0
    while i < len(smiles):
        ch = smiles[i]
        if ch == '[':
            end = smiles.index(']', i); m = re.match(r'^\d*([A-Za-z])', smiles[i + 1:end])
            if m and 'a' <= m.group(1) <= 'z': n += 1
            i = end
        elif ch in 'bcnops': n += 1
        elif (ch == 'C' and smiles[i + 1:i + 2] == 'l') or (ch == 'B' and smiles[i + 1:i + 2] == 'r'): i += 1
        i += 1
    return n


def order_key(smiles):
    from rdkit import Chem
    from rdkit.Chem import Descriptors
    m = Chem.MolFromSmiles(smiles)
    k = Chem.Mol(m); Chem.Kekulize(k, clearAromaticFlags=True)
    cc = collections.Counter()
    for b in k.GetBonds():
        if b.GetBeginAtom().GetSymbol() == 'C' and b.GetEndAtom().GetSymbol() == 'C':
            cc[b.GetBondTypeAsDouble()] += 1
    carbons = sum(1 for a in m.GetAtoms() if a.GetSymbol() == 'C')
    return (carbons, aromatic_atoms(smiles), cc[1.0], cc[2.0], cc[3.0], Descriptors.MolWt(m))


def ordered(smiles_list):
    keyed = [(order_key(s), i, s) for i, s in enumerate(smiles_list)]
    keyed.sort(key=lambda x: (-x[0][0], -x[0][1], -x[0][2], -x[0][3], -x[0][4], -x[0][5], x[1]))
    return [s for _, _, s in keyed]


def page_readable(smiles):
    """would js/chem.js parseSmiles accept it? (its element table, its lowercase rule, its valence check)"""
    from rdkit import Chem
    stripped = re.sub(r'\[[^\]]*\]', '', smiles).replace('Cl', '').replace('Br', '')
    if re.search('[a-z]', re.sub('[bcnops]', '', stripped)): return False
    m = Chem.MolFromSmiles(smiles)
    if m is None: return False
    try:
        k = Chem.Mol(m); Chem.Kekulize(k, clearAromaticFlags=True)
    except Exception:
        return False
    for a in k.GetAtoms():
        el, c = a.GetSymbol(), a.GetFormalCharge()
        if el not in VALENCE: return False
        eff = VALENCE[el] - c if el in ATE else (VALENCE[el] - abs(c) if el == 'C' else VALENCE[el] + c)
        used = sum(b.GetBondTypeAsDouble() for b in a.GetBonds())
        if used > max(eff, EXPANDED.get(el, 0)) + 1e-9: return False
    return True


# ---------- checking one recorded reaction (runs in the worker processes) ----------
_W = {}


def _init(solv, frags, roles=None):
    from rdkit import Chem, RDLogger
    RDLogger.DisableLog('rdApp.*')
    _W['roles'] = roles if roles is not None else reagent_roles()
    _W['solv'] = solv
    _W['frags'] = [(fid, Chem.MolFromSmarts(sm)) for fid, _, sm in frags]
    _W['rxn'] = {}


def _canon(smi):
    from rdkit import Chem
    m = Chem.MolFromSmiles(smi)
    return Chem.MolToSmiles(m) if m else None


def inchikey(smi):
    from rdkit import Chem
    m = Chem.MolFromSmiles(smi) if smi else None
    if m is None: return ''
    try: return Chem.MolToInchiKey(m) or ''
    except Exception: return ''


def reagent_parts(reagent):
    """the distinct parts of a reagent, sorted (a record lists each species once: K2CO3 is O=C([O-])[O-].[K+]; a
    student writes both K+; both give the same parts)"""
    return sorted(set(x for x in (reagent or '').split('.') if x))


SINGLE_ION = re.compile(r'\[[A-Z][a-z]?[+-]\d*\]')


def reagent_counted(reagent):
    """the reagent's parts as they count for a code: every molecule as many times as written (O.O = two waters, a
    different reaction from O), but a single-atom ion once (records list Na+, K+, Cl- once whatever the salt)"""
    parts = [x for x in (reagent or '').split('.') if x]
    return sorted([x for x in parts if not SINGLE_ION.fullmatch(x)] + sorted(set(x for x in parts if SINGLE_ION.fullmatch(x))))


def reagent_key(reagent):
    """the look-up key of a reagent: InChIKey of its counted parts (reagent_counted); '' = no reagent; '?...' = a
    reagent RDKit cannot key (never matches anything, so it is never mistaken for 'no reagent')"""
    if not reagent: return ''
    k = inchikey('.'.join(reagent_counted(reagent)))
    return k or '?' + reagent


LONE_ION = re.compile(r'\[[A-Za-z]+[+-]\d*\]')


def essential_parts(reagent):
    """the parts of a reagent that do the job: lone ions (Na+, K+, Cl-) left out unless nothing else is there"""
    parts = reagent_parts(reagent)
    keep = [x for x in parts if not LONE_ION.fullmatch(x)]
    return keep or parts


def subset_keys(keys, limit=8):
    """every combination of a reagent's essential part keys (the page's subsetKeys), most parts first; with more than
    `limit` parts only the combinations missing at most two of them"""
    import itertools
    keys = sorted(set(k for k in keys if k))
    low = 1 if len(keys) <= limit else len(keys) - 2
    for size in range(len(keys), low - 1, -1):
        for combo in itertools.combinations(keys, size):
            yield ' '.join(combo), size


def lookup_key(subs, reagent):
    """the key of a recorded reaction: every distinct molecule, starting materials and reagent parts together, so a
    reagent written in the starting-material boxes (or the other way round) still finds it"""
    s = sorted(set(subs))
    parts = sorted(s + [x for x in reagent_counted(reagent) if x not in s])
    k = inchikey('.'.join(parts))
    return k or '?' + '.'.join(parts)


# reagent parts that can take part in a template although they contain carbon (their atoms end up in products):
# Reagents sheet role or second role in this set. Bases such as Et3N or carbonate are left out (they would turn up as
# false nucleophiles). The page's poolExtras uses the same rule.
TRANSFER_ROLES = {'nucleophile', 'electrophile / alkylating agent', 'organometallic reagent', 'halogenating agent',
                  'protecting-group reagent', 'reactant (inorganic)', 'nucleophile (sulfur ylide precursor)'}


def pool_extras(reagent, roles):
    """as many copies of each part as the reagent lists: a reaction that needs two waters is written O.O and may use
    two; with one O written, it may use one"""
    out = []
    for x in sorted(y for y in (reagent or '').split('.') if y):
        r = roles.get(x)
        if not has_carbon(x) or (r and (r.get('role') in TRANSFER_ROLES or r.get('also') in TRANSFER_ROLES)): out.append(x)
    return out


def reagent_roles():
    p = os.path.join(HERE, 'data', 'reagents.json')
    return {x['smiles']: x for x in json.load(open(p, encoding='utf-8'))} if os.path.exists(p) else {}


def has_carbon(smi):
    from rdkit import Chem
    m = Chem.MolFromSmiles(smi)
    return bool(m) and any(a.GetSymbol() == 'C' for a in m.GetAtoms())


def _pattern_fp(m):
    from rdkit import Chem
    return Chem.PatternFingerprint(m, fpSize=2048)


def pool_fps(subs, extras):
    """pattern fingerprints of a template run's molecules (compute once per reaction, pass as run_template(screen=...))"""
    from rdkit import Chem
    return [_pattern_fp(Chem.MolFromSmiles(s)) for s in list(subs) + [x for x in extras if x not in subs]]


def could_fit(rx, fps):
    """quick screen: every reactant pattern of the template must be able to match one of the molecules (RDKit pattern
    fingerprints never miss a real substructure match, so a template that fails here could not have reacted anyway)"""
    from rdkit import DataStructs
    key = id(rx)
    q = _W.setdefault('qfp', {}).get(key)
    if q is None:
        q = [_pattern_fp(rx.GetReactantTemplate(i)) for i in range(rx.GetNumReactantTemplates())]
        _W['qfp'][key] = q
    return all(any(DataStructs.AllProbeBitsMatch(qf, mf) for mf in fps) for qf in q)


def run_template(smarts, subs, extras=(), want=None, detail=False, screen=False):
    """every distinct outcome of a template: all orders, all sites, as the page runs a learned template (runLearned).
    The molecules are the starting materials plus the reagent parts that can take part (pool_extras: carbon-free
    pieces such as OH-, Br2, HNO3, H2O, and carbon-containing nucleophiles/electrophiles such as MeO- or MeI); every
    outcome must use at least one starting material. With want (a product), also returns the leaving pieces of
    the first outcome that makes it; with detail, also which atoms of each pool molecule changed in it
    ({pool index: set of atom indices})."""
    from rdkit import Chem
    from rdkit.Chem import AllChem
    rx = _W['rxn'].get(smarts)
    if rx is None:
        try: rx = AllChem.ReactionFromSmarts(smarts); rx.Initialize()
        except Exception: rx = False
        _W['rxn'][smarts] = rx
    if rx is False: return ([], None, {}) if detail else ([], None)
    n, outs, seen, leaving, changed = rx.GetNumReactantTemplates(), [], set(), None, {}
    pool = list(subs) + [x for x in extras if x not in subs]
    mols = [Chem.MolFromSmiles(s) for s in pool]
    if screen is not False and screen is not None and not could_fit(rx, screen if isinstance(screen, list) else pool_fps(subs, extras)): return ([], None, {}) if detail else ([], None)
    for order in itertools.permutations(range(len(pool)), n):
        if min(order) >= len(subs): continue             # the reagent alone is not a reaction of these molecules
        try: res = rx.RunReactants(tuple(mols[i] for i in order), 50)
        except Exception: continue
        for prods in res:
            cs = []
            for pm in prods:
                try: cs.append(_canon(Chem.MolToSmiles(pm)))
                except Exception: cs.append(None)
            if not cs or any(c is None for c in cs): continue
            key = '.'.join(sorted(cs))
            if key not in seen: seen.add(key); outs.append(cs)
            if want is not None and leaving is None and want in cs:
                leaving = leaving_pieces([mols[i] for i in order], prods[cs.index(want)])
                if detail:
                    ch = changed_atoms([mols[i] for i in order], prods[cs.index(want)])
                    changed = {order[k]: atoms for k, atoms in ch.items()}
    return (outs, leaving, changed) if detail else (outs, leaving)


def changed_atoms(reactants, product):
    """{reactant position: atom indices that changed}: an atom changed if it left (is not in the product), its formal
    charge changed, it gained or lost a neighbour, or a non-aromatic bond to it changed order (C=C -> C-C, C=O -> C-OH).
    Hydrogen-only changes (a deprotonation) and aromatic/Kekule spelling differences do not count."""
    from rdkit import Chem
    prov = {}
    for a in product.GetAtoms():
        if a.HasProp('react_idx'): prov[a.GetIdx()] = (a.GetIntProp('react_idx'), a.GetIntProp('react_atom_idx'))
    used = collections.defaultdict(set)
    for k, i in prov.values(): used[k].add(i)
    changed = collections.defaultdict(set)
    for k, m in enumerate(reactants):
        for a in m.GetAtoms():
            if a.GetIdx() not in used[k]: changed[k].add(a.GetIdx())
    order_of = lambda b: 1.5 if (b.GetIsAromatic() or b.GetBondType() == Chem.BondType.AROMATIC) else b.GetBondTypeAsDouble()
    for p_idx, (k, i) in prov.items():
        pa, ra = product.GetAtomWithIdx(p_idx), reactants[k].GetAtomWithIdx(i)
        pn = {prov.get(b.GetOtherAtomIdx(p_idx), ('new', b.GetOtherAtomIdx(p_idx))): order_of(b) for b in pa.GetBonds()}
        rn = {(k, b.GetOtherAtomIdx(i)): order_of(b) for b in ra.GetBonds()}
        if set(pn) != set(rn) or pa.GetFormalCharge() != ra.GetFormalCharge():
            changed[k].add(i); continue
        if any(pn[key] != rn[key] and 1.5 not in (pn[key], rn[key]) for key in pn):
            changed[k].add(i)
    return changed


def reacted_fragments(smiles, atoms):
    """the fragments of one starting material that reacted: an atom of one of its matches changed"""
    from rdkit import Chem
    if not atoms: return []
    m = Chem.MolFromSmiles(smiles)
    out = []
    for fid, q in _W['frags']:
        if q is not None and any(atoms.intersection(match) for match in m.GetSubstructMatches(q, uniquify=True, maxMatches=1000)):
            out.append(fid)
    return sorted(out)


def leaving_pieces(reactants, product):
    """the reactant atoms that are not in the product, as molecules (H added where a bond to the product broke)"""
    from rdkit import Chem
    used = collections.defaultdict(set)
    for a in product.GetAtoms():
        d = a.GetPropsAsDict()
        if 'react_idx' in d: used[d['react_idx']].add(d['react_atom_idx'])
    pieces = []
    for k, m in enumerate(reactants):
        gone = [a.GetIdx() for a in m.GetAtoms() if a.GetIdx() not in used[k]]
        if not gone or len(gone) == m.GetNumAtoms(): continue          # all kept, or a whole spectator
        rw = Chem.RWMol(m)
        for idx in sorted(used[k], reverse=True): rw.RemoveAtom(idx)
        for a in rw.GetAtoms():                              # where a bond broke, an H takes its place (not on ions or metals)
            if a.GetFormalCharge() == 0 and a.GetSymbol() not in ATE - {'B'}:
                a.SetNoImplicit(False)
        try:
            frag = rw.GetMol(); Chem.SanitizeMol(frag)
            pieces += Chem.MolToSmiles(frag).split('.')
        except Exception:
            pieces.append('?')
    return '.'.join(sorted(pieces))


def frag_counts(smiles_list):
    from rdkit import Chem
    c = collections.Counter()
    for s in smiles_list:
        m = Chem.MolFromSmiles(s)
        for fid, q in _W['frags']:
            if q is not None:
                n = len(m.GetSubstructMatches(q, uniquify=True, maxMatches=1000))
                if n: c[fid] += n
    return c


def check(task):
    i, r, p, a, smarts = task
    rec = [_canon(x) for x in r.split('.')]
    prod = _canon(p)
    if not prod or any(s is None for s in rec): return (i, 'unreadable SMILES')
    # carbon-free reactants (Br2, HNO3, OH-, NH3, H2O, N3- ...) are written as the reagent, as in reaction_rules.xlsx
    subs = [x for x in rec if has_carbon(x)]
    if not subs: return (i, 'no carbon-containing starting material')
    if prod in subs: return (i, 'product is a starting material')
    comps = [x for x in a.split('.') if x] if a else []
    conds, reag = set(), [x for x in rec if x not in subs]
    for x in comps:
        cx = _canon(x)
        if cx is None: return (i, 'unreadable SMILES')
        if cx in _W['solv']: conds.add(_W['solv'][cx])
        else: reag.append(cx)
    reagent = _canon('.'.join(reag)) if reag else ''
    if reag and not reagent: return (i, 'unreadable SMILES')
    extras = pool_extras(reagent, _W['roles'])
    outs, leaving, changed = run_template(smarts, subs, extras, want=prod, detail=True)
    if not any(prod in o for o in outs): return (i, 'template does not give the recorded product')
    reacted = {s: reacted_fragments(s, changed.get(j, set())) for j, s in enumerate(subs)}   # pool index j = subs[j]
    a0, b0 = frag_counts(subs), frag_counts([prod])
    consumed = sorted(f for f in a0 if b0[f] < a0[f])
    formed = sorted(f for f in b0 if b0[f] > a0[f])
    subs = ordered(subs)
    reacted_col = ' | '.join(' + '.join(reacted.get(x) or []) or '-' for x in subs)
    readable = all(page_readable(x) for x in subs + [prod] + ([reagent] if reagent else []))
    parts = sorted(set(inchikey(x) for x in reagent_parts(reagent)))
    return (i, 'ok', subs, reagent, ', '.join(sorted(conds, key=str.lower)), prod, len(outs), consumed, formed,
            leaving or '', readable, lookup_key(subs, reagent), reagent_key(reagent), ' '.join(p for p in parts if p), reacted_col)


# ---------- order ----------
def candidates(K):
    """every recorded reaction, in round-robin template order (commonest template first; a new reagent first)"""
    tn, tid = K['tn'], K['tid']
    by_t = collections.defaultdict(list)
    for i in range(len(K['r'])):
        if '*' in K['r'][i] or '*' in K['p'][i]: continue
        by_t[tid[i]].append(i)
    order = sorted(by_t, key=lambda t: (-tn[t], t))
    lists = []
    for t in order:
        idx = sorted(by_t[t], key=lambda i: (not K['yl'][i], len(K['p'][i]) + len(K['r'][i]), i))
        firsts, rest, seen = [], [], set()
        for i in idx:
            rg = K['a'][i]
            (rest if rg in seen else firsts).append(i); seen.add(rg)
        lists.append(firsts + rest)
    out = []
    for k in itertools.count():
        got = False
        for L in lists:
            if k < len(L): out.append(L[k]); got = True
        if not got: break
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--n', type=int, default=0, help='learn only this many (0 = all)')
    ap.add_argument('--write-only', action='store_true', help='only write the workbook again from the last run (if saving failed)')
    args = ap.parse_args()
    t0 = time.time()
    K = load_kb()
    if args.write_only:
        import pickle
        chosen, why_not = pickle.load(open(SAVED, 'rb'))
        print(f'loaded {len(chosen):,} learned reactions from the last run ({SAVED})')
        write_workbook(K, chosen, why_not)
        return
    solv, frags = solvents(), fragments()
    print(f'knowledge base: {len(K["r"]):,} recorded reactions, {len(K["tmpl"]):,} templates ({time.time() - t0:.0f} s)')
    order = candidates(K)
    why_not = collections.Counter({'no structure (*) in the record': len(K['r']) - len(order)})
    chosen, by_code = [], {}
    want = args.n or len(order)
    with Pool(max(1, (os.cpu_count() or 2) - 1), initializer=_init, initargs=(solv, frags, reagent_roles())) as pool:
        pos, step = 0, 50000
        while len(chosen) < want and pos < len(order):
            batch = order[pos:pos + step]; pos += step
            tasks = [(i, K['r'][i], K['p'][i], K['a'][i], K['tmpl'][K['tid'][i]]) for i in batch]
            res = {x[0]: x for x in pool.imap_unordered(check, tasks, chunksize=250)}
            for i in batch:                                # keep the round-robin order
                x = res[i]
                if x[1] != 'ok': why_not[x[1]] += 1; continue
                code = '~'.join(['.'.join(x[2]), x[3], x[4], x[5]])
                if code in by_code: chosen[by_code[code]][1].append(i); why_not['same reaction recorded again (counted)'] += 1; continue
                if len(chosen) >= want: continue
                by_code[code] = len(chosen); chosen.append([x, [i]])
            print(f'  checked {min(pos, len(order)):,} of {len(order):,}: {len(chosen):,} learned ({time.time() - t0:.0f} s)', flush=True)
    print('not learned:', dict(why_not))
    json.dump({'checked': min(pos, len(order)) + why_not['no structure (*) in the record'], 'learned': len(chosen), 'why_not': why_not},
              open(os.path.join(HERE, 'learned_summary.json'), 'w', encoding='utf-8'), indent=1)
    import pickle
    pickle.dump((chosen, why_not), open(SAVED, 'wb'))     # kept, so a failed save is retried in minutes (--write-only)
    write_workbook(K, chosen, why_not)
    print(f'{len(chosen):,} learned reactions ({time.time() - t0:.0f} s)')


def write_workbook(K, chosen, why_not):
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Alignment
    F = 'Arial'
    HDR, BODY, FILL = Font(name=F, bold=True, color='FFFFFF'), Font(name=F, size=10), PatternFill('solid', fgColor='1F4E78')
    wb = Workbook(write_only=True)
    col = lambda k: (chr(64 + k // 26) if k >= 26 else '') + chr(65 + k % 26)

    def sheet(title, head, widths, rows):
        ws = wb.create_sheet(title)
        for k, w in enumerate(widths): ws.column_dimensions[col(k)].width = w
        ws.freeze_panes = 'A2'
        cells = []
        for h in head:
            c = WriteOnlyCell(ws, value=h); c.font, c.fill = HDR, FILL; c.alignment = Alignment(wrap_text=True, vertical='center'); cells.append(c)
        ws.append(cells)
        for r in rows:
            ws.append(r)                                   # plain values: 1 M styled cells would make the file huge

    tn, tid = K['tn'], K['tid']
    per_t = collections.defaultdict(list)
    rows = []
    num = lambda v: float(v) if re.fullmatch(r'\d+(\.\d+)?', str(v or '')) else None
    for n, (x, idxs) in enumerate(chosen, 1):
        i, _, subs, reagent, conds, prod, nout, consumed, formed, leaving, readable, skey, rkey, pkeys, reacted_col = x
        t = tid[i]; lid = f'L{n:07d}' if len(chosen) >= 1000000 else f'L{n:06d}'
        per_t[t].append((lid, consumed, formed, leaving or ''))
        yields = [num(K['yl'][j]) for j in idxs if num(K['yl'][j]) is not None]
        rows.append([lid, '~'.join(['.'.join(subs), reagent, conds, prod]), '.'.join(subs), reagent, conds, prod, f'T{t + 1:06d}', tn[t],
                     ' + '.join(consumed), ' + '.join(formed), reacted_col, leaving, nout, len(idxs), K['c'][i],
                     int(K['y'][i]) if str(K['y'][i]).isdigit() else K['y'][i], max(yields) if yields else None,
                     'yes', 'yes' if readable else 'no', skey, rkey, pkeys])
    head = ['Learned ID', 'Reaction code (starting materials~reagent~conditions~product)', 'Starting materials', 'Reagent',
            'Conditions (recorded solvents)', 'Product', 'Template ID', 'Times this template is recorded', 'Fragments consumed',
            'Fragments formed', 'Fragments that reacted, per starting material (in code order, | between molecules)', 'Leaving pieces (starting-material atoms not in the product)', 'Outcomes of the template on these molecules',
            'Times this reaction is recorded', 'Source (patent / ORD id)', 'Year', 'Best recorded yield (%)',
            'Checked (RDKit: template gives the recorded product)', 'Page can read every molecule',
            'Look-up key (InChIKey of starting materials + reagent parts)', 'Reagent key (InChIKey of its distinct parts)', 'Reagent part keys']
    widths = [10, 70, 40, 30, 22, 40, 11, 12, 18, 18, 24, 24, 12, 12, 16, 7, 8, 14, 10, 30, 30, 30]
    for k in range(0, max(1, len(rows)), ROWS_PER_SHEET):
        sheet(f'Learned {k // ROWS_PER_SHEET + 1}', head, widths, rows[k:k + ROWS_PER_SHEET])
    trows = []
    for t in sorted(per_t, key=lambda t: (-tn[t], t)):
        ex = per_t[t]
        cons = collections.Counter(' + '.join(e[1]) for e in ex).most_common(1)[0][0]
        form = collections.Counter(' + '.join(e[2]) for e in ex).most_common(1)[0][0]
        leave = collections.Counter(e[3] for e in ex).most_common(1)[0][0]
        trows.append([f'T{t + 1:06d}', K['tmpl'][t], tn[t], len(ex), cons, form, leave, ex[0][0]])
    sheet('Templates', ['Template ID', 'Reaction SMARTS (rdchiral template: the atoms that change + neighbours)', 'Times recorded (all 1.1 M reactions)',
                        'Learned reactions', 'Fragments consumed (most common)', 'Fragments formed (most common)',
                        'Leaving pieces (most common)', 'First example'],
          [11, 90, 14, 10, 20, 20, 24, 12], trows)
    meta = K['meta']
    nsheets = (len(rows) - 1) // ROWS_PER_SHEET + 1
    sheet('How to use', ['How this workbook is organised'], [140], [[t] for t in [
        f'Learned 1-{nsheets}: {len(rows):,} learned reactions, one per row ({ROWS_PER_SHEET:,} rows per sheet). Same four-part code as reaction_rules.xlsx.',
        f'Templates: the {len(trows):,} different reaction templates they use (the reaction SMARTS the page runs on new molecules).',
        'Not learned (with the reason): ' + '; '.join(f'{k}: {v:,}' for k, v in why_not.items()) + '.',
        'The hand-made rules (R001...) in reaction_rules.xlsx come first because they explain (fragments, conditions, pKa, stereo, mechanisms). '
        'When a rule answers AND these exact molecules + reagent are a recorded reaction, the page compares them: same product -> the rule, citing the '
        'record; the record adds stereo the rule lacks -> the recorded stereo; a different product -> the recorded product, saying what the rule would give '
        'and that the record\'s conditions were not compared. A rule result that depends on conditions or amounts (no reaction without heat, excess, 2 equiv) '
        'is never replaced, only annotated. When no rule answers, the learned reactions are looked up, then predicted.',
        'Same starting materials + reagent as a learned row: the page gives that row\'s recorded product (looked up). New molecules: the page runs the '
        'templates recorded with that reagent and ranks the products they give (predicted), and says which.',
        'Conditions are the recorded solvents (Solvents sheet of reaction_rules.xlsx). Recorded reactions have no temperature, so none is given.',
        'Fragments consumed / formed: read with the Fragments sheet of molecule_data.xlsx (count drops / rises from the starting materials to the product).',
        'Leaving pieces: the starting-material atoms the product does not contain (H added where a bond broke): what by-products are made from.',
        'Fragments that reacted: for each starting material (same order as the code), the fragments with an atom that made, broke or changed a bond '
        '(found by following every atom from the starting materials into the product); "-" = that molecule\'s fragments were all left untouched.',
        'No mechanisms: recorded reactions give starting materials -> product only.',
        'Checked: RDKit ran each row\'s template on its starting materials (every order and site, as the page does) and got the recorded product.',
        f'Source: {meta.get("source", "")}',
        f'Licence: {meta.get("license", "")}',
        'Rebuild: python tools/learn_reactions.py, then python tools/learned_to_js.py; re-check: python tools/check_learned.py.']])
    save_workbook(wb, OUT)


if __name__ == '__main__':
    main()
