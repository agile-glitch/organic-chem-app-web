#!/usr/bin/env python3
"""Give the AGILES spreadsheet's molecules the names of the IUPAC namer as it is now (js/chem.js with
data/ring_parents.js), without remaking the codes: each molecule is named from its own code's drawing, each name is read
back by OPSIN (the open-source IUPAC name reader) and compared with the molecule (RDKit, full InChIKey), and the code's
AGILES.iupac and AGILES.iupac_checked items and the IUPAC name / IUPAC name checked columns are rewritten. Nothing
else in a code changes (its drawing, identifiers and labels are as they were); a Name that was the old IUPAC name follows
the new one. A molecule the namer cannot name keeps its old name. IUPAC name checked: yes when OPSIN reads the name
back to exactly the molecule, stereo included, else no.

    python tools/agiles_rename.py            AGILES_codes.xlsx and AGILES_codes.csv, in place
Needs Node.js, Java + OPSIN (as tools/namer_bench/bench.py) and RDKit. tools/agiles_table.js remakes the whole
spreadsheet from the website instead (Playwright + Chromium).
"""
import base64, csv, json, os, subprocess, sys, zlib
import openpyxl
from rdkit import Chem, RDLogger

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(HERE, 'namer_bench'))
from bench import opsin                                    # noqa: E402
RDLogger.DisableLog('rdApp.*')
MAX_HEAVY = 100                                            # as the Molecule tab (NAMER_MAX_HEAVY): larger ones are not named


def unpack(code):
    b = code.split(':', 1)[1]
    return zlib.decompress(base64.urlsafe_b64decode(b + '=' * (-len(b) % 4)), -15).decode('utf-8')


def pack(sdf):
    c = zlib.compressobj(9, zlib.DEFLATED, -15)
    return 'AGILES1:' + base64.urlsafe_b64encode(c.compress(sdf.encode('utf-8')) + c.flush()).decode('ascii').rstrip('=')


def set_item(sdf, key, value):
    """the record with item AGILES.<key> set to value (removed when value is None); a new item goes before smiles"""
    lines = sdf.split('\n')
    head = f'>  <AGILES.{key}>'
    if head in lines:
        i = lines.index(head); j = i + 1
        while j < len(lines) and lines[j] != '': j += 1
        lines[i:j + 1] = [] if value is None else [head, value, '']
    elif value is not None:
        for anchor in ('>  <AGILES.iupac_checked>', '>  <AGILES.smiles>', '>  <AGILES.smiles_rdkit>', '>  <AGILES.inchi>', '>  <AGILES.formula>'):
            if anchor in lines and not (key == 'iupac_checked' and anchor == '>  <AGILES.iupac_checked>'):
                k = lines.index(anchor)
                if key == 'iupac_checked' and '>  <AGILES.iupac>' in lines: k = lines.index('>  <AGILES.iupac>') + 3
                lines[k:k] = [head, value, '']
                break
    return '\n'.join(lines)


def inline_strings(path):
    """rewrite the workbook with each text in its own cell (inline strings), as tools/agiles_table.py writes it:
    readers of the spreadsheet (AgileSynth's Lookup) take the text from the cell, not from a shared-strings table"""
    import html, re, shutil, zipfile
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    tmp = path + '.tmp'
    with zipfile.ZipFile(path) as zin:
        shared = []
        if 'xl/sharedStrings.xml' in zin.namelist():
            import xml.etree.ElementTree as ET
            for si in ET.fromstring(zin.read('xl/sharedStrings.xml')).iter(f'{{{ns}}}si'):
                shared.append(''.join(t.text or '' for t in si.iter(f'{{{ns}}}t')))
        def cell(m):
            attrs, idx = m.group(1), int(m.group(2))
            text = html.escape(shared[idx], quote=False)
            sp = ' xml:space="preserve"' if shared[idx] != shared[idx].strip() else ''
            return f'<c{attrs.replace(' t="s"', ' t="inlineStr"')}><is><t{sp}>{text}</t></is></c>'
        with zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item)
                if item.filename.startswith('xl/worksheets/sheet'):
                    data = re.sub(r'<c([^>]*? t="s"[^>]*)><v>(\d+)</v></c>', cell, data.decode('utf-8')).encode('utf-8')
                elif item.filename == 'xl/sharedStrings.xml':
                    continue
                elif item.filename == '[Content_Types].xml':
                    data = re.sub(rb'<Override PartName="/xl/sharedStrings.xml"[^>]*/>', b'', data)
                elif item.filename == 'xl/_rels/workbook.xml.rels':
                    data = re.sub(rb'<Relationship [^>]*sharedStrings[^>]*/>', b'', data)
                zout.writestr(item, data)
    shutil.move(tmp, path)


def main():
    xlsx, csvp = os.path.join(APP, 'AGILES_codes.xlsx'), os.path.join(APP, 'AGILES_codes.csv')
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]
    rows = [r for r in ws.iter_rows(min_row=3) if r[1].value and str(r[1].value).startswith('AGILES1:')]
    print(f'{len(rows):,} molecules', flush=True)
    sdfs = [unpack(r[1].value) for r in rows]
    mbs = [s.split('M  END', 1)[0] + 'M  END\n' for s in sdfs]
    mols = [Chem.MolFromMolBlock(mb) for mb in mbs]
    # the namer, on each code's drawing
    req = '\n'.join(json.dumps({'id': i, 'molfile': mb}) for i, mb in enumerate(mbs) if mols[i] is not None and mols[i].GetNumHeavyAtoms() <= MAX_HEAVY) + '\n'
    p = subprocess.run(['node', os.path.join(HERE, 'namer_bench', 'namer.js')], input=req, capture_output=True, text=True, encoding='utf-8')
    names = {}
    for line in p.stdout.splitlines():
        d = json.loads(line)
        if d.get('name'): names[d['id']] = d['name']
    print(f'{len(names):,} named by the namer', flush=True)
    # every name (new, or the old one kept) read back by OPSIN
    final = [names.get(i) or (r[2].value or None) for i, r in enumerate(rows)]
    ids = [i for i, n in enumerate(final) if n]
    smis = opsin([final[i] for i in ids])
    ok = {}
    for i, smi in zip(ids, smis):
        back = Chem.MolFromSmiles(smi) if smi else None
        ok[i] = bool(back is not None and mols[i] is not None and Chem.MolToInchiKey(back) == Chem.MolToInchiKey(mols[i]))
    changed = kept = 0
    for i, r in enumerate(rows):
        old = r[2].value or None
        new = final[i]
        chk = ('yes' if ok.get(i) else 'no') if new else None
        if new != old: changed += 1
        if i not in names and old: kept += 1
        sdf = set_item(set_item(sdfs[i], 'iupac', new), 'iupac_checked', chk)
        code = pack(sdf) if sdf != sdfs[i] else r[1].value
        if r[0].value == old and new: r[0].value = new
        r[1].value, r[2].value, r[3].value = code, new, chk
        # the code still holds the same drawing
        back = Chem.MolFromMolBlock(unpack(code).split('M  END', 1)[0] + 'M  END\n')
        if mols[i] is not None and (back is None or Chem.MolToInchiKey(back) != Chem.MolToInchiKey(mols[i])):
            raise SystemExit(f'row {r[0].row}: the rewritten code is not the same drawing')
    # the About sheet's line on the names
    about = wb.worksheets[1]
    for row in about.iter_rows():
        c = row[0]
        if c.value and str(c.value).startswith('IUPAC name: '):
            c.value = ('IUPAC name: the app\'s namer (js/chem.js with data/ring_parents.js; tools/agiles_rename.py); "yes" in the '
                       'checked column: OPSIN, the open-source IUPAC name reader, reads the name back to exactly this molecule, '
                       'stereo included; "no": it does not (treat such a name with care). A blank cell: no name.')
    wb.save(xlsx)
    inline_strings(xlsx)
    with open(csvp, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow([c.value for c in ws[2]])
        for r in rows: w.writerow(['' if c.value is None else c.value for c in r])
    yes = sum(1 for i in range(len(rows)) if ok.get(i))
    print(f'{changed:,} names changed ({kept:,} old names kept where the namer gives none); '
          f'{yes:,} of {len(rows):,} read back by OPSIN to exactly the molecule; '
          f'{sum(1 for n in final if not n):,} without a name -> {os.path.relpath(xlsx, APP)}, {os.path.relpath(csvp, APP)}')


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
    main()
