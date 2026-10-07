"""Learn which reaction wins when several could happen (chemoselectivity / regioselectivity), and measure it.

    python tools/learn_selectivity.py                 # learn win rates, measure on held-out reactions, write the page's weights
    python tools/learn_selectivity.py --train 40000 --test 3000

The problem: for new molecules the page runs every template learned with the reagent. Often several fit (an amine
AND an alcohol that could both be acylated; two different ring positions; a ketone and an ester for a hydride) and
they give different products. Which one really happens?

What is learned (from the recorded reactions, no hand-written rules): for each recorded reaction, every template
learned with its reagent is run on its starting materials, exactly as the page does. A template that fits is
"applicable"; if it gives the recorded product it "won", otherwise it "lost" (something else happened instead). A
template's win rate = wins / times applicable (smoothed: (wins + 1) / (applicable + 2)). A template with a high win
rate describes a site or group that reacts first; a low one describes a site that is usually beaten by another.

How it is measured (honestly): the learned reactions are split by Learned ID; 1 in 20 is held out (TEST) and never
used to learn anything (not the look-up index, not the win rates). Each held-out reaction is then predicted from its
starting materials + reagent alone, and the prediction counts as right only if it is the recorded product. Several
ways of ranking the products are compared, and the best one is what the page uses.

Writes data/learned/w/*.js (the win rates, for the page) and tools/selectivity_report.json (the measurements).
"""
import argparse, collections, hashlib, itertools, json, os, pickle, sys, tempfile, time
from multiprocessing import Pool

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
sys.path.insert(0, HERE)
SRC = os.path.join(APP, 'learned_reactions.xlsx')
CACHE = os.path.join(tempfile.gettempdir(), 'organic_chem_learned_rows.pkl')
TOP = None                                             # templates tried per reaction: all of them, as on the page
MIN_OBS = 2                                            # ...that were seen at least twice in that stage (the page's MIN_OBSERVATIONS)


def is_test(lid):
    return int(hashlib.md5(lid.encode()).hexdigest(), 16) % 20 == 0


def read_rows():
    """[Learned ID, starting materials list, reagent, product, template #, reagent key, part keys] + the template SMARTS"""
    if os.path.exists(CACHE) and os.path.getmtime(CACHE) > os.path.getmtime(SRC):
        return pickle.load(open(CACHE, 'rb'))
    import openpyxl
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    trows = [r for r in list(wb['Templates'].iter_rows(values_only=True))[1:] if r and r[0]]
    tindex = {r[0]: k for k, r in enumerate(trows)}
    rows = []
    for name in wb.sheetnames:
        if not name.startswith('Learned'): continue
        head = None
        for r in wb[name].iter_rows(values_only=True):
            if head is None: head = {h: k for k, h in enumerate(r)}; continue
            if not r or not r[0]: continue
            g = lambda col: r[head[col]]
            rows.append([g('Learned ID'), (g('Starting materials') or '').split('.'), g('Reagent') or '', g('Product'), tindex[g('Template ID')],
                         g('Reagent key (InChIKey of its distinct parts)') or '', (g('Reagent part keys') or '').split(), g('Page can read every molecule') == 'yes'])
    data = {'rows': rows, 'smarts': [r[1] for r in trows], 'tid': [r[0] for r in trows]}
    pickle.dump(data, open(CACHE, 'wb'))
    return data


def build_index(rows):
    by_reagent = collections.defaultdict(collections.Counter)
    by_set = collections.defaultdict(collections.Counter)
    import learn_reactions as L
    ess = {}
    for r in rows:
        if r[2] and not r[5].startswith('?'): by_reagent[r[5]][r[4]] += 1
        if r[2]:
            if r[2] not in ess:
                ks = sorted(L.inchikey(x) for x in L.essential_parts(r[2]))
                ess[r[2]] = ' '.join(ks) if all(ks) else ''
            if ess[r[2]]: by_set[ess[r[2]]][r[4]] += 1
    return ({k: c.most_common() for k, c in by_reagent.items()}, {k: c.most_common() for k, c in by_set.items()})


def exact_list(r, idx):
    """templates learned with this reagent, plus with this reagent + a starting material that is itself a known reagent
    (MeI typed as a starting material: records often list it WITH the reagents); counts summed, most seen first, only
    those seen at least MIN_OBS times (the page's first stage does the same)"""
    import learn_reactions as L
    by_reagent, roles = idx[0], L._W.get('roles', {})
    keys = [] if r[5].startswith('?') else [r[5]]
    for s in r[1]:
        if s in roles:
            k = L.reagent_key(r[2] + '.' + s if r[2] else s)
            if k and not k.startswith('?'): keys.append(k)
    sums = {}
    for k in dict.fromkeys(keys):
        for t, n in by_reagent.get(k, []): sums[t] = sums.get(t, 0) + n
    return sorted([(t, n) for t, n in sums.items() if n >= MIN_OBS], key=lambda x: -x[1])


def candidates(r, idx):
    """the page's candidate list: templates learned with this reagent, else with its parts (lone ions ignored)"""
    by_reagent, by_part = idx
    reag = r[2]
    if not reag: return [], 'none'                       # the page does not predict without a reagent (only looks up)
    lst = exact_list(r, idx)
    if lst: return lst[:TOP], 'reagent'
    # parts: only records whose reagent is PART of this one (every combination of its essential parts)
    from learn_reactions import inchikey, essential_parts, subset_keys
    keys = [inchikey(p) for p in essential_parts(reag)]
    sums = {}
    for combo, size in subset_keys(keys):
        for t, n in by_part.get(combo, []):
            o = sums.setdefault(t, [0, 0]); o[0] += n; o[1] = max(o[1], size)
    lst = sorted(sums.items(), key=lambda kv: (-kv[1][1], -kv[1][0]))
    return [(t, v[0]) for t, v in lst if v[0] >= MIN_OBS][:TOP], 'parts'


_G = {}


def _init(smarts, idx, solv_frags):
    import learn_reactions as L
    L._init(*solv_frags, L.reagent_roles())
    _G['smarts'], _G['idx'] = smarts, idx


def outcomes_for(r):
    """run every candidate template on a reaction's starting materials: {template: [outcome keys]}"""
    import learn_reactions as L
    lst, how = candidates(r, _G['idx'])
    extras = L.pool_extras(r[2], L._W['roles'])
    fps = L.pool_fps(r[1], extras)                     # fingerprints once per reaction: the quick screen per template
    res = []
    for t, n in lst:
        outs, _ = L.run_template(_G['smarts'][t], r[1], extras, screen=fps)
        if outs: res.append((t, n, ['.'.join(sorted(o)) for o in outs]))
    return r[0], r[3], how, res


def rank(res, method, win):
    """products ranked by a method; res = [(template, count, [outcome keys])]"""
    score = collections.defaultdict(float)
    for t, n, outs in res:
        w = win.get(t, (0, 0)); wr = (w[0] + 1) / (w[1] + 2)
        for o in dict.fromkeys(outs):                        # in outcome order, as the page (ties break the same way)
            if method == 'count': score[o] += n                       # what the page did first: templates weighted by count
            elif method == 'winrate-max': score[o] = max(score[o], wr + 1e-6 * n)
            elif method == 'count*winrate': score[o] += n * wr
            elif method == 'winrate-sum': score[o] += wr
            elif method == 'count*winrate^2': score[o] += n * wr * wr
    return sorted(score, key=lambda o: -score[o])


def addition_kind(smarts):
    """'C=C', 'C#C', 'C=O', 'C=N' when the template adds across that pi bond (a mapped bond whose order drops and both
    atoms stay), else ''. Read from the template itself."""
    from rdkit import Chem
    try:
        left, right = smarts.split('>>')
        L, R = Chem.MolFromSmarts(left), Chem.MolFromSmarts(right)
    except Exception:
        return ''
    if L is None or R is None: return ''
    rmap = {a.GetAtomMapNum(): a.GetIdx() for a in R.GetAtoms() if a.GetAtomMapNum()}
    for b in L.GetBonds():
        a1, a2 = b.GetBeginAtom(), b.GetEndAtom()
        m1, m2 = a1.GetAtomMapNum(), a2.GetAtomMapNum()
        if not (m1 and m2 and m1 in rmap and m2 in rmap): continue
        o = b.GetBondType()
        if o not in (Chem.BondType.DOUBLE, Chem.BondType.TRIPLE) or b.GetIsAromatic(): continue
        rb = R.GetBondBetweenAtoms(rmap[m1], rmap[m2])
        if rb is None: continue
        ro = rb.GetBondTypeAsDouble()
        if ro < b.GetBondTypeAsDouble() and not rb.GetIsAromatic():
            els = ''.join(sorted([a1.GetSymbol(), a2.GetSymbol()]))
            if els == 'CC': return 'C#C' if o == Chem.BondType.TRIPLE else 'C=C'
            if els == 'CO': return 'C=O'
            if els == 'CN': return 'C=N' if o == Chem.BondType.DOUBLE else 'C#N'
    return ''


METHODS = ['count', 'count*winrate', 'count*winrate^2', 'winrate-sum', 'winrate-max']


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--train', type=int, default=60000, help='training reactions used to learn the win rates')
    ap.add_argument('--test', type=int, default=4000, help='held-out reactions to measure on')
    ap.add_argument('--workers', type=int, default=0, help='worker processes (each holds a copy of the index; fewer = less memory)')
    args = ap.parse_args()
    import learn_reactions as L
    t0 = time.time()
    D = read_rows()
    rows = D['rows']
    train = [r for r in rows if not is_test(r[0])]
    test = [r for r in rows if is_test(r[0])]
    print(f'{len(rows):,} learned reactions: {len(train):,} to learn from, {len(test):,} held out ({time.time() - t0:.0f} s)', flush=True)
    idx = build_index(train)
    solv_frags = (L.solvents(), L.fragments())
    step_tr = max(1, len(train) // args.train)
    step_te = max(1, len(test) // args.test)
    sample_tr, sample_te = train[::step_tr][:args.train], test[::step_te][:args.test]
    win = collections.defaultdict(lambda: [0, 0])
    with Pool(args.workers or max(1, (os.cpu_count() or 2) - 1), initializer=_init, initargs=(D['smarts'], idx, solv_frags)) as pool:
        done = 0
        for lid, prod, how, res in pool.imap_unordered(outcomes_for, sample_tr, chunksize=20):
            for t, n, outs in res:
                w = win[t]; w[1] += 1
                if prod in outs: w[0] += 1
            done += 1
            if done % 5000 == 0: print(f'  learned from {done:,} of {len(sample_tr):,} ({time.time() - t0:.0f} s)', flush=True)
        win = {t: tuple(v) for t, v in win.items()}
        hits = {m: [0, 0, 0] for m in METHODS}              # top-1, top-3, any
        add_hits = collections.defaultdict(lambda: {m: [0, 0] for m in METHODS})     # addition reactions: [right, tested]
        kind_of = {r[0]: addition_kind(D['smarts'][r[4]]) for r in sample_te}
        by_how = collections.Counter(); n_te = 0; examples = []
        for lid, prod, how, res in pool.imap_unordered(outcomes_for, sample_te, chunksize=20):
            n_te += 1; by_how[how] += 1
            kind = kind_of.get(lid, '')
            for m in METHODS:
                rk = rank(res, m, win)
                if rk[:1] and prod in rk[0].split('.'): hits[m][0] += 1
                if kind:
                    add_hits[kind][m][1] += 1; add_hits['all additions'][m][1] += 1
                    if rk[:1] and prod in rk[0].split('.'): add_hits[kind][m][0] += 1; add_hits['all additions'][m][0] += 1
                if any(prod in o.split('.') for o in rk[:3]): hits[m][1] += 1
                if any(prod in o.split('.') for o in rk): hits[m][2] += 1
            if len(examples) < 30:
                rk = rank(res, 'count', win); rk2 = rank(res, 'count*winrate', win)
                examples.append({'id': lid, 'recorded': prod, 'count': rk[:1], 'count*winrate': rk2[:1]})
    report = {'held_out_tested': n_te, 'learned_from': len(sample_tr), 'candidates_from': dict(by_how),
              'methods': {m: {'top1': round(h[0] / n_te, 4), 'top3': round(h[1] / n_te, 4), 'product_among_any': round(h[2] / n_te, 4)} for m, h in hits.items()},
              'additions': {k: {m: {'top1': round(v[m][0] / max(1, v[m][1]), 4), 'tested': v[m][1]} for m in METHODS} for k, v in add_hits.items()},
              'examples': examples}
    best = max(METHODS, key=lambda m: hits[m][0])
    report['best'] = best
    print(json.dumps(report['methods'], indent=1)); print('best:', best)
    print('addition reactions (top-1 with the best method):', {k: v[best] for k, v in report['additions'].items()})
    json.dump(report, open(os.path.join(HERE, 'selectivity_report.json'), 'w', encoding='utf-8'), indent=1)
    # the page's weights: [wins, times applicable] per template number
    out = os.path.join(APP, 'data', 'learned', 'w'); os.makedirs(out, exist_ok=True)
    from learned_to_js import packed
    with open(os.path.join(out, 'winrates.js'), 'w', encoding='utf-8') as f:
        f.write('(window.LEARNED_PARTS=window.LEARNED_PARTS||{})["w/winrates"]={"gz":"' +
                packed({'method': best, 'win': {str(t): v for t, v in win.items()}}) + '"};\n')
    print(f'wrote data/learned/w/winrates.js ({len(win):,} templates) and tools/selectivity_report.json ({time.time() - t0:.0f} s)')


if __name__ == '__main__':
    main()
