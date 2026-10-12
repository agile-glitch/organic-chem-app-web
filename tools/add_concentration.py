"""Teach the workbook about concentration: two new conditions, 'concentrated' and 'dilute', and the rules that depend on them.

    python tools/add_concentration.py        # idempotent; then python tools/rules_to_js.py

The Reactions page lets the user right-click the reagent and set its strength (37 % HCl, 98 % H2SO4, 3 % H2O2 ...). js/reagent_conc.js turns
that strength into one of these two conditions for the reagent in question (what counts as concentrated differs: 37 % is conc. HCl, 98 % is conc.
H2SO4, 50 % is conc. NaOH). The words can also be typed in the conditions box ("conc.", "dilute").

How a rule uses them (js/reactions.js requirementMet): a rule that lists 'concentrated' or 'dilute' under Required conditions is satisfied
  * when the strength has not been given at all (the textbook conditions of the rule are assumed, exactly as before this column existed),
  * or when it has been given and falls in that class. A strength given and in the other class (or in neither: 50 % H2SO4 is not conc.)
    stops the rule, and the page says why: "needs concentrated; the reagent is set to 10 %".
Chemistry behind each one (the textbook condition each rule is already named for in its Conditions cell):
"""
import os, sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
PATH = os.path.join(os.path.dirname(HERE), 'reaction_rules.xlsx')

CONDITIONS = [
    ('concentrated', 'concentrated, conc, conc., concd',
     "The reagent is used at full strength (conc. H2SO4 ~98 %, conc. HCl ~37 %, conc. HNO3 ~70 %, conc. NaOH ~50 %). Set from the reagent's "
     "strength (right-click the reagent) or typed here. A rule that requires it is assumed to have it when no strength is given."),
    ('dilute', 'dilute, diluted, dil, dil.',
     "The reagent is used as a dilute aqueous solution (about 1-2 M, 10 % or weaker). Water is then the nucleophile and the medium, so addition/hydrolysis "
     "wins over dehydration/elimination. Set from the reagent's strength (right-click the reagent) or typed here."),
]
# rule id -> the class to add to its Required conditions (the Conditions cell of each already says so in words)
RULES = {
    'R006': 'dilute',        # acid-catalysed hydration of an alkene: dilute aqueous H2SO4 (conc. acid + heat dehydrates instead, R035)
    'R018': 'concentrated',  # nitration: conc. HNO3 + conc. H2SO4 make the nitronium ion
    'R035': 'concentrated',  # dehydration of an alcohol: conc. H2SO4 (or H3PO4), 160 C
    'R050': 'dilute',        # acid-catalysed epoxide opening: dilute aqueous H2SO4
    'R065': 'concentrated',  # Clemmensen: Zn(Hg) in conc. HCl
    'R080': 'concentrated',  # Cannizzaro: conc. NaOH
    'R089': 'dilute',        # acetal hydrolysis: dilute aqueous acid, excess water
    'R093': 'concentrated',  # Beckmann: conc. H2SO4
    'R097': 'dilute',        # acid-catalysed ester hydrolysis: dilute acid, excess water
}


def main():
    wb = load_workbook(PATH)
    ws = wb['Conditions']
    head = [c.value for c in ws[1]]
    have = {str(r[0].value).strip().lower() for r in ws.iter_rows(min_row=2) if r[0].value}
    added = 0
    for name, words, note in CONDITIONS:
        if name in have: continue
        row = [None] * len(head)
        row[0] = name; row[1] = words
        row[head.index('Notes')] = note
        ws.append(row); added += 1
    wr = wb['Reactions']
    h = [c.value for c in wr[1]]
    ci = next(i for i, v in enumerate(h) if str(v).startswith('Required conditions'))
    changed = []
    for r in wr.iter_rows(min_row=2):
        rid = str(r[0].value or '').strip()
        if rid in RULES:
            cell = r[ci]
            cur = str(cell.value or '').strip()
            if RULES[rid] in [p.strip().lower() for p in cur.replace('|', ',').split(',')]: continue
            if '|' in cur:
                print(f'{rid}: has alternatives ({cur!r}); not changed', file=sys.stderr); continue
            cell.value = (cur + ', ' if cur else '') + RULES[rid]
            changed.append(rid)
    if added or changed:
        wb.save(PATH)
    print(f'conditions added: {added}; rules changed: {", ".join(changed) or "none"}')


if __name__ == '__main__':
    main()
