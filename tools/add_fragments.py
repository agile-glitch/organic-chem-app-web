"""Add the learned fragments to the Fragments sheet of molecule_data.xlsx.

    python tools/add_fragments.py --dry-run        # print the rows, change nothing
    python tools/add_fragments.py                  # append the new groups of tools/new_fragments.py (text: tools/new_fragments_text.py)
    python tools/add_fragments.py --sub            # append the sub-fragments of tools/new_subfragments.py (text: tools/new_subfragments_text.py)
    (then run tools/rules_to_js.py)

New groups: run tools/discover_fragments.py and tools/fragment_eval.py first; the evidence in each Notes cell (how many starting materials carry the
group, how often it took part in the recorded reaction; how many learned reactions use the reagent) comes from tools/fragment_candidates_scored.json.
Sub-fragments: run tools/subfragment_scan.py, then tools/subfragment_eval.py; the evidence (how often the parent group is consumed with and without the
feature) comes from tools/subfragment_scored.json.
Idempotent: a candidate whose SMARTS is already on the sheet is skipped, so a second run adds nothing. A copy of the workbook as it was is kept in the temp folder.
"""
import argparse, copy, json, os, re, shutil, sys, tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
sys.path.insert(0, HERE)
import new_fragments as NF
import new_fragments_text as NT

XLSX = os.path.join(APP, 'molecule_data.xlsx')
SCORED = os.path.join(HERE, 'fragment_candidates_scored.json')
SUB_SCORED = os.path.join(HERE, 'subfragment_scored.json')
DROP = {'aziridine'}                                # too rare in the recorded reactions to describe
SAMPLE = 70570                                      # starting materials in the sample tools/discover_fragments.py scored them on
ALL = 1090843                                       # learned reactions
PKA_SOURCE = 'typical values in water (DMSO where marked): Evans pKa table; Bordwell; Clayden, Organic Chemistry (approximate: the rest of the molecule shifts them)'
SOURCE = 'learned from the recorded reactions (tools/discover_fragments.py, tools/new_fragments.py; RDKit-checked SMARTS)'
SUB_SOURCE = 'sub-fragment learned from the recorded reactions (tools/subfragment_scan.py, tools/new_subfragments.py; RDKit-checked SMARTS)'


def cell(v):
    return None if v in ('None', '', None) else v


def evidence(c, sc):
    if not sc: return ''
    if c['kind'] == 'reagent':
        return f"[learned from the recorded reactions: a part of the reagent in {sc['reagent_reactions']:,} of the {ALL:,} learned reactions]"
    s, r = sc['support'], sc['reacted']
    return (f"[learned from the recorded reactions: found in {s / SAMPLE:.1%} of {SAMPLE:,} sampled starting materials; an atom of it took part in the "
            f"recorded reaction in {r / max(s, 1):.0%} of those ({r:,} of {s:,}); {sc['new']:,} of those reaction centres had no specific fragment on the sheet before]")


def sub_evidence(c, sc):
    if not sc or sc.get('rate_with') is None: return ''
    gain = f"; on reactions the rates had not seen it adds {sc['gain_held_nats']:,.0f} nats of log-likelihood" if sc.get('gain_held_nats') is not None else ''
    return (f"[learned from the recorded reactions: in a sample of 200,000 training reactions {sc['support_train']:,} have this group together with its parent; the parent group was "
            f"consumed in {sc['rate_with']:.0%} of them against {sc['rate_without']:.0%} of the {sc['without_train']:,} that have the parent without this feature, for the same kind of reagent{gain}]")


def candidates_new_groups(scored):
    out = []
    for c in NF.CANDIDATES:
        if c['slug'] in DROP: continue
        sc = scored.get(c['slug'])
        out.append({'slug': c['slug'], 'name': c['name'], 'smarts': c['smarts'], 'text': NT.TEXT[c['slug']], 'evidence': evidence(c, sc), 'source': SOURCE})
    return out


def candidates_sub(scored):
    import new_subfragments as NS, new_subfragments_text as NST
    out = []
    for d in NS.SUBFRAGMENTS:
        sc = scored.get(d['slug'])
        out.append({'slug': d['slug'], 'name': d['name'], 'smarts': d['smarts'], 'text': NST.TEXT[d['slug']], 'evidence': sub_evidence(d, sc), 'source': SUB_SOURCE})
    return out


def main():
    ap = argparse.ArgumentParser(); ap.add_argument('--dry-run', action='store_true'); ap.add_argument('--sub', action='store_true', help='add the sub-fragments (tools/new_subfragments.py)')
    a = ap.parse_args()
    import openpyxl
    path = SUB_SCORED if a.sub else SCORED
    scored = {x['slug']: x for x in json.load(open(path, encoding='utf-8'))}
    cands = candidates_sub(scored) if a.sub else candidates_new_groups(scored)
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['Fragments']
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    last = max(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value)
    ids = [str(ws.cell(r, 1).value).strip() for r in range(2, last + 1) if ws.cell(r, 1).value]
    have_smarts = {str(ws.cell(r, col['SMARTS pattern'] + 1).value).strip() for r in range(2, last + 1)}
    nxt = max(int(re.sub(r'\D', '', i)) for i in ids if re.fullmatch(r'F\d+', i)) + 1
    added, manifest = [], {}
    for c in cands:
        if c['smarts'] in have_smarts: continue
        t = c['text']
        fid = f'F{nxt:03d}'; nxt += 1
        notes = (t['notes'].rstrip('. ') + '. ' + c['evidence']).strip()
        row = {'Fragment ID': fid, 'Fragment name': c['name'], 'SMARTS pattern': c['smarts'], 'Reactive atom(s)': t['reactive'], 'Role': t['role'],
               'pKa (its most acidic H)': cell(t['acid']), 'pKa of its conjugate acid (how basic)': cell(t['base']), 'Leaving group': cell(t['lg']),
               'Made into a good leaving group by': None, 'Steric hindrance (at the reactive atom)': t['steric'], 'Ring strain': t.get('ring', 'none: not a ring fragment'),
               'pKa source': PKA_SOURCE, 'Source': c['source'], 'Notes': notes}
        added.append(row); manifest[c['slug']] = fid
    print(f'{len(added)} fragments to add ({len(ids)} on the sheet, next id {added[0]["Fragment ID"] if added else "-"})')
    if a.dry_run:
        for r in added[:6]: print(json.dumps(r, ensure_ascii=False)[:700])
        return
    if not added: return
    shutil.copy(XLSX, os.path.join(tempfile.gettempdir(), 'molecule_data_before_' + ('subfragments' if a.sub else 'new_fragments') + '.xlsx'))
    style_row = last
    for k, row in enumerate(added):
        r = last + 1 + k
        for h, i in col.items():
            src = ws.cell(style_row, i + 1); dst = ws.cell(r, i + 1)
            dst.value = row.get(h)
            dst._style = copy.copy(src._style)
        ws.row_dimensions[r].height = ws.row_dimensions[style_row].height
    wb.save(XLSX)
    mpath = os.path.join(HERE, 'data', 'fragments_new_ids.json')
    old = json.load(open(mpath, encoding='utf-8')) if os.path.exists(mpath) else {}
    old.update(manifest)
    json.dump(old, open(mpath, 'w', encoding='utf-8'), indent=1)
    print('appended', ', '.join(f'{v} {k}' for k, v in manifest.items()))
    print('now run: python tools/rules_to_js.py')


if __name__ == '__main__':
    main()
