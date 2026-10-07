"""Turn learned_reactions.xlsx into data/learned/ for the Reactions page (a web page cannot read an Excel file).

    python tools/learned_to_js.py

Up to a million+ learned reactions are far too much for a page to load at once, so the data is cut into small files
and the page loads only the ones a reaction needs (gzip inside base64, as data/rxnkb, so a double-clicked page can
read them too):
  learned_meta.js   counts and chunk sizes
  rx/NNNNN.js       the learned reactions, RX_CHUNK per file, in Learned-ID order:
                    [Learned ID, code, template #, consumed, formed, leaving pieces, times recorded, source, year, yield]
  t/NNN.js          the templates (reaction SMARTS), TMPL_CHUNK per file, commonest first: {id, smarts, recorded}
  x/AB.js           look-up: InChIKey of every distinct molecule (starting materials + reagent parts) -> learned reaction
                    numbers (AB = its first two letters); the page's learnedLookupKey computes the same
  g/AB.js           reagent key -> the templates learned with it: [[template #, how many learned reactions, example #], ...]
                    (AB = first two letters of the reagent InChIKey; '__' = no reagent)
  s/AB.js           a recorded reagent's essential parts (InChIKeys, sorted, space-joined) -> the same. The page looks up
                    every combination of YOUR reagent's essential parts, so it only borrows from records whose reagent
                    is part of yours (Br2 + HBr is never borrowed for HBr alone)
  r/XX.js           reagent role signature (XX = role_shard: sum of its character codes mod 256, in hex) (Reagents sheet roles, e.g. "base" or "base|coupling reagent") -> the same
                    (the last resort: reagents that do the same job, K2CO3 ~ Cs2CO3)
Keys are RDKit InChIKeys, exactly what the page's keyOf gives. Also writes tools/learned_tests.json (a sample of
learned reactions for the page's "Run all tests" button).
"""
import base64, collections, gzip, json, os, shutil, time

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
SRC = os.path.join(APP, 'learned_reactions.xlsx')
OUT = os.path.join(APP, 'data', 'learned')
RX_CHUNK, TMPL_CHUNK, MAX_PER_REAGENT = 5000, 5000, None    # None: every template learned with a reagent is kept
LOOKUP_COL = 'Look-up key (InChIKey of starting materials + reagent parts)'
REAGENT_COL = 'Reagent key (InChIKey of its distinct parts)'


def packed(obj):
    return base64.b64encode(gzip.compress(json.dumps(obj, separators=(',', ':'), ensure_ascii=False).encode('utf-8'), 9)).decode('ascii')


SKIP_ROLES = {'counter-ion / spectator ion', 'solvent', 'salt (work-up or additive)', 'other', 'drying agent'}


def role_signature(reagent, roles):
    """the jobs a reagent's parts do (Reagents sheet), spectators and solvents left out; None if a part is unknown
    (the page's roleSignature computes the same)"""
    sig = set()
    for p in reagent.split('.'):
        r = roles.get(p)
        if r is None: return None
        if r['role'] not in SKIP_ROLES: sig.add(r['role'])
    return '|'.join(sorted(sig)) if sig else None


def role_shard(sig):
    """the r/ file of a role signature (the page's roleShard computes the same)"""
    return '%02x' % (sum(ord(c) for c in sig) % 256)


def shard(key):
    return key[:2] if key else '__'


def main():
    import openpyxl
    t0 = time.time()
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    tall = list(wb['Templates'].iter_rows(values_only=True))
    thead = {h: k for k, h in enumerate(tall[0])}
    trows = [r for r in tall[1:] if r and r[0]]
    tleave = [(r[thead['Leaving pieces (most common)']] or '') if 'Leaving pieces (most common)' in thead else '' for r in trows]
    tindex = {r[0]: k for k, r in enumerate(trows)}
    rows, readable, exact = [], [], collections.defaultdict(list)
    roles = {x['smiles']: x for x in json.load(open(os.path.join(HERE, 'data', 'reagents.json'), encoding='utf-8'))}         if os.path.exists(os.path.join(HERE, 'data', 'reagents.json')) else {}
    by_role = collections.defaultdict(collections.Counter)
    by_reagent = collections.defaultdict(collections.Counter)
    by_set = collections.defaultdict(collections.Counter)
    import learn_reactions as L
    ess_key = {}
    example = {}
    for name in wb.sheetnames:
        if not name.startswith('Learned'): continue
        head = None
        for r in wb[name].iter_rows(values_only=True):
            if head is None: head = {h: k for k, h in enumerate(r)}; continue
            if not r or not r[0]: continue
            g = lambda col: r[head[col]]
            k = len(rows)
            code = g('Reaction code (starting materials~reagent~conditions~product)')
            t = tindex[g('Template ID')]
            yl = g('Best recorded yield (%)')
            rows.append([g('Learned ID'), code, t, g('Fragments consumed') or '', g('Fragments formed') or '',
                         g('Leaving pieces (starting-material atoms not in the product)') or '', g('Times this reaction is recorded') or 1,
                         g('Source (patent / ORD id)') or '', g('Year') or '', yl if yl is not None else ''])
            readable.append(g('Page can read every molecule') == 'yes')
            skey, rkey, pkeys = g(LOOKUP_COL) or '', g(REAGENT_COL) or '', (g('Reagent part keys') or '').split()
            exact[skey].append(k)
            if not rkey.startswith('?'):                       # a reagent RDKit cannot key is never filed as 'no reagent'
                by_reagent[rkey][t] += 1
                example.setdefault((rkey, t), k)
            reagent = code.split('~')[1]
            if reagent:
                if reagent not in ess_key:
                    ks = sorted(L.inchikey(x) for x in L.essential_parts(reagent))
                    ess_key[reagent] = ' '.join(ks) if all(ks) else ''
                sk = ess_key[reagent]
                if sk: by_set[sk][t] += 1; example.setdefault(('s', sk, t), k)
            sig = role_signature(reagent, roles) if reagent else None
            if sig: by_role[sig][t] += 1; example.setdefault(('r', sig, t), k)
            if len(rows) % 100000 == 0: print(f'  read {len(rows):,} ({time.time() - t0:.0f} s)', flush=True)
    print(f'read {len(rows):,} learned reactions, {len(trows):,} templates ({time.time() - t0:.0f} s)')

    for d in ('rx', 't', 'x', 'g', 'p', 's', 'r'):                    # (w/ holds the selectivity weights: tools/learn_selectivity.py)
        if os.path.isdir(os.path.join(OUT, d)): shutil.rmtree(os.path.join(OUT, d))
        os.makedirs(os.path.join(OUT, d))

    def write(name, obj):
        with open(os.path.join(OUT, name + '.js'), 'w', encoding='utf-8') as f:
            f.write(f'(window.LEARNED_PARTS=window.LEARNED_PARTS||{{}})[{json.dumps(name)}]={{"gz":"{packed(obj)}"}};\n')

    for k in range(0, len(rows), RX_CHUNK): write(f'rx/{k // RX_CHUNK:05d}', rows[k:k + RX_CHUNK])
    for k in range(0, len(trows), TMPL_CHUNK):
        part = trows[k:k + TMPL_CHUNK]
        write(f't/{k // TMPL_CHUNK:03d}', {'id': [r[0] for r in part], 'smarts': [r[1] for r in part], 'recorded': [r[2] for r in part],
                                          'leaving': tleave[k:k + TMPL_CHUNK]})
    shards = collections.defaultdict(dict)
    for key, ks in exact.items():
        if not key.startswith('?'): shards[shard(key)][key] = ks
    for s, d in shards.items(): write(f'x/{s}', d)
    nx = len(shards)
    shards = collections.defaultdict(dict)
    for rk, cnt in by_reagent.items(): shards[shard(rk)][rk] = [[t, n, example[(rk, t)]] for t, n in cnt.most_common(MAX_PER_REAGENT)]
    for s, d in shards.items(): write(f'g/{s}', d)
    shards = collections.defaultdict(dict)
    for sk, cnt in by_set.items(): shards[shard(sk)][sk] = [[t, n, example[('s', sk, t)]] for t, n in cnt.most_common(MAX_PER_REAGENT)]
    for s, d in shards.items(): write(f's/{s}', d)
    shards = collections.defaultdict(dict)
    for sig, cnt in by_role.items(): shards[role_shard(sig)][sig] = [[t, n, example[('r', sig, t)]] for t, n in cnt.most_common(MAX_PER_REAGENT)]
    for s_, d in shards.items(): write(f'r/{s_}', d)
    kbm = json.loads(open(os.path.join(APP, 'data', 'rxnkb', 'meta.js'), encoding='utf-8').read().split('=', 1)[1].rstrip(';\n '))
    meta = {'n': len(rows), 'n_templates': len(trows), 'n_reagents': len(by_reagent), 'n_keys': len(exact), 'rx_chunk': RX_CHUNK,
            'tmpl_chunk': TMPL_CHUNK, 'source': kbm.get('source', ''), 'license': kbm.get('license', ''), 'built': time.strftime('%Y-%m-%d')}
    with open(os.path.join(OUT, 'learned_meta.js'), 'w', encoding='utf-8') as f:
        f.write('window.LEARNED_META=' + json.dumps(meta) + ';\n')
    size = sum(os.path.getsize(os.path.join(dp, x)) for dp, _, fs in os.walk(OUT) for x in fs)
    nfiles = sum(len(fs) for _, _, fs in os.walk(OUT))
    print(f'wrote {OUT}: {len(rows):,} reactions, {len(trows):,} templates, {len(by_reagent):,} reagents, {len(exact):,} look-up keys '
          f'in {nx} shards; {nfiles:,} files, {size / 1e6:.0f} MB ({time.time() - t0:.0f} s)')

    # a sample for the page's "Run all tests" button: evenly spread learned reactions the page can read, looked up from their codes
    step = max(1, len(rows) // 50)
    tests = [{'feature': 'learned reactions (looked up)', 'name': f'{r[0]} (template {trows[r[2]][0]})',
              'code': '~'.join(r[1].split('~')[:3]) + '~', 'expect': r[0], 'products': [r[1].split('~')[3]]}
             for k, r in enumerate(rows) if k % step == 0 and readable[k]][:50]
    json.dump({'learned': tests}, open(os.path.join(HERE, 'learned_tests.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'wrote tools/learned_tests.json: {len(tests)} learned reactions for the Run all tests button')


if __name__ == '__main__':
    main()
