"""Check every reaction in reaction_rules.xlsx against tools/reaction_tests.json, without the browser.

Layer 1  SMARTS:      the reaction's SMARTS, run by RDKit on its example molecules, can give the expected product.
Layer 2  fragments:   the fragments the reaction says it consumes are in the starting materials, and the ones it
                      says it forms are in the products (fragment SMARTS from molecule_data.xlsx).
Layer 3  mechanism:   every step's curved arrows, replayed with electron bookkeeping, give that step's result
                      (tools/check_mechanisms.py), and the last step contains the product.
Layer 4  the page:    run in the browser (the same JSON file), not here.

    python tools/run_all_checks.py
"""
import json, os, re, sys
from itertools import permutations
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rdkit import Chem, RDLogger
from rdkit.Chem import AllChem
from openpyxl import load_workbook
from check_mechanisms import statuses, canon_set
RDLogger.DisableLog("rdApp.*")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

def canon(s, stereo=True):
    m = Chem.MolFromSmiles(s)
    if m is None: return None
    if not stereo: Chem.RemoveStereochemistry(m)
    return Chem.MolToSmiles(m)

def sheet(path, name):
    ws = load_workbook(os.path.join(ROOT, path), read_only=True)[name]
    it = ws.iter_rows(values_only=True); head = [str(h) for h in next(it)]
    return [dict(zip(head, r)) for r in it if r and r[0]]
find = lambda row, start: next((v for k, v in row.items() if k.startswith(start)), None)

tests = json.load(open(os.path.join(ROOT, "tools", "reaction_tests.json"), encoding="utf-8"))
rx = {str(r["Reaction ID"]): r for r in sheet("reaction_rules.xlsx", "Reactions")}
frag = {str(f["Fragment ID"]): Chem.MolFromSmarts(f["SMARTS pattern"]) for f in sheet("molecule_data.xlsx", "Fragments") if f.get("SMARTS pattern")}
mech_rows = {}
for m in sheet("reaction_rules.xlsx", "Mechanism"):
    if isinstance(m.get("Step"), (int, float)):
        mech_rows.setdefault(str(m["Reaction ID"]), []).append(m)
mstat, _ = statuses()

# fragment ID rule: a descriptive ID (F_NAME) per structurally distinct SMARTS; no two fragments share an ID or a pattern
def check_fragment_ids():
    rows = sheet("molecule_data.xlsx", "Fragments"); bad = []
    ids, pats = {}, {}
    for f in rows:
        fid, sm = str(f["Fragment ID"]).strip(), str(f.get("SMARTS pattern") or "").strip()
        if not re.fullmatch(r"F_[A-Z0-9]+(?:_[A-Z0-9]+)*", fid): bad.append("bad ID format: " + fid)
        if fid in ids: bad.append("duplicate ID: " + fid)
        ids[fid] = sm
        q = Chem.MolFromSmarts(sm) if sm else None
        if q is None: bad.append(fid + ": missing or invalid SMARTS"); continue
        key = Chem.MolToSmarts(q)
        if key in pats: bad.append(f"{fid} and {pats[key]} have the same SMARTS")
        pats.setdefault(key, fid)
    return bad
_bad_ids = check_fragment_ids()
if _bad_ids:
    print("FRAGMENT ID CHECK FAILED:\n  " + "\n  ".join(_bad_ids)); sys.exit(1)

def items(text):
    return [re.findall(r"F_[A-Z0-9]+(?:_[A-Z0-9]+)*", part) for part in str(text or "").split("+") if re.findall(r"F_[A-Z0-9]+(?:_[A-Z0-9]+)*", part)]

rows, fails, nomechs = [], 0, 0
for t in tests["reactions"]:
    rid, r = t["id"], rx.get(t["id"])
    res = {"id": rid}
    if r is None:
        res.update(smarts="MISSING", frags="-", mech="-"); rows.append(res); fails += 1; continue
    # layer 1: SMARTS on the example molecules (every site, every order), expected product among the outcomes
    want = {canon(p, False) for p in t["products"]}
    rxn = AllChem.ReactionFromSmarts(find(r, "Reaction SMARTS"))
    n = rxn.GetNumReactantTemplates(); seen = set()
    for order in permutations(range(len(t["substrates"])), n):
        for out in rxn.RunReactants(tuple(Chem.MolFromSmiles(t["substrates"][i]) for i in order), 50):
            try:
                ps = set()
                for p in out:
                    Chem.SanitizeMol(p); ps.add(canon(Chem.MolToSmiles(p), False))
                seen.add(frozenset(ps))
            except Exception:
                pass
    ok1 = any(want <= s for s in seen)
    res["smarts"] = "pass" if ok1 else "FAIL"
    # layer 2: fragments consumed / formed
    subs = [Chem.MolFromSmiles(s) for s in t["substrates"]]; prods = [Chem.MolFromSmiles(s) for s in t["products"]]
    has = lambda mols, ids: any(m.HasSubstructMatch(frag[i]) for m in mols for i in ids if i in frag)
    bad2 = [("consumes " + "|".join(a)) for a in items(find(r, "Fragments consumed")) if not has(subs, a)] + \
           [("forms " + "|".join(a)) for a in items(find(r, "Fragment formed")) if not has(prods, a)]
    res["frags"] = "pass" if not bad2 else "FAIL (" + "; ".join(bad2) + ")"
    # layer 3: mechanism replay + the last step contains the product
    st = (mstat or {}).get(rid, {}).get("status", "none")
    steps = sorted(mech_rows.get(rid, []), key=lambda m: m["Step"])
    last_ok = bool(steps) and want <= {canon(x, False) for x in str(find(steps[-1], "Intermediate") or "").split(".") if canon(x, False)}
    nomech = st == "none" and not steps            # mechanisms are optional: products only, arrows learned later
    res["mech"] = "not yet (optional)" if nomech else (st if st.startswith("verified") else "FAIL " + st) + ("" if last_ok or not steps else "; last step lacks the product")
    if not (ok1 and not bad2 and (nomech or (st.startswith("verified") and last_ok))):
        fails += 1
    nomechs += nomech
    rows.append(res)

print(f"{'reaction':9s} {'1 SMARTS':9s} {'2 fragments':12s} 3 mechanism (replay)")
for r in rows:
    print(f"{r['id']:9s} {r['smarts']:9s} {r['frags'][:12]:12s} {r['mech']}")
print(f"\n{len(rows) - fails} of {len(rows)} reactions pass layers 1-3" + ("" if not fails else f"; {fails} need attention")
      + (f" ({nomechs} have no mechanism yet, so layer 3 is skipped for them)" if nomechs else ""))
